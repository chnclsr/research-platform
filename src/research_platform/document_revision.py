"""Plan, render, validate and publish immutable revisions of generated reports."""

from __future__ import annotations

import copy
import hashlib
import io
import json
import re
import zipfile
from collections import Counter
from dataclasses import fields
from typing import Any

from docx import Document
from docx.shared import Pt as DocxPt
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.util import Cm as PptxCm
from pptx.util import Pt as PptxPt

from .config import get_settings
from .figure_analysis import FigureObservation, GeneratedResearchFigure
from .formula_render import load_formula_displays
from .language_guard import foreign_sentences, language_matches
from .llm import LLMProvider
from .presentation_polisher import polish_and_record
from .presentation_report import build_presentation_report
from .report_synthesis import (
    StudyProfile,
    SynthesisPackage,
    SynthesisSection,
    citation_tokens,
)
from .repository import Repository, RevisionConflict
from .schemas import (
    CoverageMetrics,
    ReportCitation,
    ResearchProtocol,
    RevisionPlan,
    RevisionStatus,
)
from .storage import ObjectStore
from .word_report import build_word_report

PLAN_PROMPT_VERSION = "document-revision-v1"
DOCX_MEDIA_TYPE = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
PPTX_MEDIA_TYPE = "application/vnd.openxmlformats-officedocument.presentationml.presentation"
VERSIONED_BUNDLES = frozenset({"result_bundle.zip", "research_bundle.zip"})
TERMINAL_RUN_STATUSES = frozenset({"completed", "completed_incomplete"})
_SECTION_TEXT_FIELDS = frozenset(
    {"title", "synthesis", "consensus", "disagreements", "implications"}
)
_PACKAGE_TEXT_FIELDS = frozenset(
    {"executive_summary", "cross_study_assessment", "conclusion", "uncertainty"}
)
_PPTX_BOUNDS_TARGET = re.compile(r"^pptx\.slide:(\d+)\.shape:(.+)\.bounds$")
_PPTX_FONT_TARGET = re.compile(r"^pptx\.slide:(\d+)\.shape:(.+)\.font_size_pt$")
_DOCX_FONT_TARGET = re.compile(r"^docx\.style:(.+)\.font_size_pt$")
_FORMAT_FEEDBACK = re.compile(
    r"\b(?:slide|slayt|pptx|powerpoint|word|docx|font|typeface|yazı|görsel|resim|"
    r"image|picture|chart|grafik|layout|düzen|boyut|büyüt|küçült)\b",
    re.IGNORECASE,
)
_LOCKED_NUMBER = re.compile(
    r"(?<![\w.])[+-]?\d+(?:[.,]\d+)?(?:\s?(?:%|mg|g|kg|ml|l|mm|cm|m|km|h|s))?",
    re.IGNORECASE,
)
_LOCKED_QUOTE = re.compile(r"(?:\"[^\"\n]+\"|“[^”\n]+”|‘[^’\n]+’)")


class RevisionValidationError(ValueError):
    """The proposed plan or its rendered output crossed a revision safety boundary."""


def _section_id(index: int, title: str) -> str:
    digest = hashlib.sha256(f"{index}:{title}".encode()).hexdigest()[:12]
    return f"section-{index + 1:02d}-{digest}"


def report_model_from_manifest(manifest: dict[str, Any]) -> dict[str, Any]:
    """Turn a shipped reproducibility manifest into the stable editable model."""
    protocol = ResearchProtocol.model_validate(manifest.get("protocol") or {})
    synthesis = copy.deepcopy(manifest.get("synthesis") or {})
    sections = []
    for index, raw in enumerate(synthesis.get("sections") or []):
        section = dict(raw)
        section["id"] = str(section.get("id") or _section_id(index, str(section.get("title", ""))))
        sections.append(section)
    synthesis["sections"] = sections
    return {
        "schema_version": 1,
        "title": protocol.title_for_report(),
        "question": protocol.question_for_report(),
        "language": protocol.report_language,
        "protocol": protocol.model_dump(mode="json"),
        "coverage": dict(manifest.get("coverage") or {}),
        "synthesis": synthesis,
    }


def synthesis_package_from_model(report_model: dict[str, Any]) -> SynthesisPackage:
    payload = copy.deepcopy(report_model.get("synthesis") or {})
    section_fields = {item.name for item in fields(SynthesisSection)}
    profile_fields = {item.name for item in fields(StudyProfile)}
    package_fields = {item.name for item in fields(SynthesisPackage)}
    payload["sections"] = [
        SynthesisSection(**{key: value for key, value in section.items() if key in section_fields})
        for section in payload.get("sections") or []
    ]
    payload["study_profiles"] = [
        StudyProfile(**{key: value for key, value in profile.items() if key in profile_fields})
        for profile in payload.get("study_profiles") or []
    ]
    return SynthesisPackage(**{key: value for key, value in payload.items() if key in package_fields})


def editable_targets(report_model: dict[str, Any]) -> dict[str, str]:
    synthesis = report_model.get("synthesis") or {}
    targets = {
        "document.title": str(report_model.get("title") or ""),
        **{name: str(synthesis.get(name) or "") for name in _PACKAGE_TEXT_FIELDS},
    }
    for section in synthesis.get("sections") or []:
        section_id = str(section.get("id") or "")
        for name in _SECTION_TEXT_FIELDS:
            targets[f"section:{section_id}.{name}"] = str(section.get(name) or "")
    return targets


def plan_context(
    report_model: dict[str, Any], format_targets: dict[str, Any] | None = None
) -> dict[str, Any]:
    """The bounded model surface the planning LLM may address."""
    targets = editable_targets(report_model)
    return {
        "language": report_model.get("language"),
        "targets": [
            {
                "id": target,
                "text": text,
                "citations": sorted(set(citation_tokens(text))),
            }
            for target, text in targets.items()
        ],
        "section_order": [
            str(section.get("id"))
            for section in (report_model.get("synthesis") or {}).get("sections") or []
        ],
        "format_targets": format_targets or {},
    }


async def generate_revision_plan(
    llm: LLMProvider,
    report_model: dict[str, Any],
    feedback: str,
    *,
    format_targets: dict[str, Any] | None = None,
) -> tuple[RevisionPlan, list[dict[str, Any]]]:
    """Ask the model for data, never for an Office file or executable instructions."""
    system = """You are a conservative research-report editor. Return one JSON object only.
Use only the target ids supplied by the caller. Never invent, remove or alter source labels
such as [S03], numeric findings, quotations or evidence. If the request needs new research,
set classification to research_extension_required and return no operations. If one short
question is required, set requires_clarification and provide it. Otherwise return:
{
  "summary": "plain-language summary",
  "classification": "document_revision",
  "requires_clarification": false,
  "clarification_question": "",
  "operations": [{
    "operation": "set_title|replace_text|rewrite_section|shorten_section|rename_section|reorder_sections|set_format_override",
    "target": "an allowed target id, sections, or pptx./docx. override key",
    "instruction": "the user's instruction",
    "rationale": "why",
    "format_scope": "all|docx|pptx",
    "preserve_citations": true,
    "replacement": "complete final text for text operations",
    "order": [],
    "value": null,
    "expected_summary": "visible result"
  }]
}
For reorder_sections, target must be sections and order must contain every section id once.
For set_format_override use only one of these exact target grammars:
- pptx.slide:<1-based number>.shape:<exact shape name>.bounds with a value object containing
  left_cm, top_cm, width_cm and/or height_cm;
- pptx.slide:<1-based number>.shape:<exact shape name>.font_size_pt with a numeric value;
- docx.style:<exact style name>.font_size_pt with a numeric value.
PPTX shape names must come from format_targets. Match a natural description to a shape only
when the slide and shape type/text make the choice unambiguous. Otherwise ask one short
clarifying question; never ask the user for an internal shape name."""
    user = json.dumps(
        {
            "feedback": feedback,
            "document": plan_context(report_model, format_targets),
        },
        ensure_ascii=False,
    )
    raw = await llm.complete_json(system, user)
    if isinstance(raw, dict) and isinstance(raw.get("plan"), dict):
        raw = raw["plan"]
    plan = RevisionPlan.model_validate(raw)
    validate_plan_targets(report_model, plan)
    return plan, llm.drain_metrics()


def validate_plan_targets(report_model: dict[str, Any], plan: RevisionPlan) -> None:
    allowed = set(editable_targets(report_model))
    section_ids = {
        str(section.get("id"))
        for section in (report_model.get("synthesis") or {}).get("sections") or []
    }
    for operation in plan.operations:
        if not operation.preserve_citations and operation.operation != "set_format_override":
            raise RevisionValidationError("Report citations are locked during document revision")
        if operation.operation == "reorder_sections":
            if operation.target != "sections" or set(operation.order) != section_ids:
                raise RevisionValidationError(
                    "reorder_sections must contain every existing section id exactly once"
                )
            if len(operation.order) != len(section_ids):
                raise RevisionValidationError("reorder_sections contains duplicate ids")
        elif operation.operation == "set_format_override":
            prefix = operation.target.split(".", 1)[0]
            if prefix not in {"docx", "pptx"}:
                raise RevisionValidationError("format override target must start with docx. or pptx.")
            if operation.format_scope not in {prefix, "all"}:
                raise RevisionValidationError("format override scope does not match its target")
            _validate_format_override(operation.target, operation.value)
        elif operation.format_scope != "all":
            raise RevisionValidationError(
                "Content operations must update the shared DOCX/PPTX report model"
            )
        elif operation.target not in allowed:
            raise RevisionValidationError(f"Unknown revision target: {operation.target}")
        if operation.operation == "set_title" and operation.target != "document.title":
            raise RevisionValidationError("set_title may target only document.title")


def _validate_format_override(target: str, value: Any) -> None:
    bounds = _PPTX_BOUNDS_TARGET.fullmatch(target)
    pptx_font = _PPTX_FONT_TARGET.fullmatch(target)
    docx_font = _DOCX_FONT_TARGET.fullmatch(target)
    if bounds:
        if int(bounds.group(1)) < 1 or not isinstance(value, dict):
            raise RevisionValidationError("PPTX bounds require a slide and an object")
        allowed = {"left_cm", "top_cm", "width_cm", "height_cm"}
        if not value or set(value) - allowed:
            raise RevisionValidationError(
                "PPTX bounds accept left_cm, top_cm, width_cm and height_cm only"
            )
        for key, raw in value.items():
            try:
                number = float(raw)
            except (TypeError, ValueError) as exc:
                raise RevisionValidationError(f"{key} must be numeric") from exc
            if number < 0 or number > 60 or (key in {"width_cm", "height_cm"} and not number):
                raise RevisionValidationError(f"{key} is outside the supported page bounds")
        return
    if pptx_font or docx_font:
        try:
            number = float(value)
        except (TypeError, ValueError) as exc:
            raise RevisionValidationError("font_size_pt must be numeric") from exc
        if not 6 <= number <= 72:
            raise RevisionValidationError("font_size_pt must be between 6 and 72")
        if pptx_font and int(pptx_font.group(1)) < 1:
            raise RevisionValidationError("PPTX slides are numbered from 1")
        return
    raise RevisionValidationError(
        "Unsupported format override; use PPTX shape bounds/font size or DOCX style font size"
    )


def _set_target(report_model: dict[str, Any], target: str, value: str) -> None:
    if target == "document.title":
        report_model["title"] = value
        return
    synthesis = report_model["synthesis"]
    if target in _PACKAGE_TEXT_FIELDS:
        synthesis[target] = value
        return
    section_target, field_name = target.split(".", 1)
    section_id = section_target.removeprefix("section:")
    for section in synthesis.get("sections") or []:
        if section.get("id") == section_id:
            section[field_name] = value
            return
    raise RevisionValidationError(f"Unknown section target: {target}")


def _validate_replacement_language(text: str, language: str) -> None:
    words = [item for item in text.split() if any(char.isalpha() for char in item)]
    if len(words) < 12:
        return
    if not language_matches(text, language) or foreign_sentences(text, language):
        raise RevisionValidationError("A revision operation changed the report language")


def apply_revision_plan(
    report_model: dict[str, Any], plan: RevisionPlan
) -> tuple[dict[str, Any], dict[str, Any]]:
    if plan.classification != "document_revision":
        raise RevisionValidationError("This request requires a new research run")
    if plan.requires_clarification:
        raise RevisionValidationError("The revision still requires clarification")
    validate_plan_targets(report_model, plan)
    result = copy.deepcopy(report_model)
    overrides: dict[str, Any] = {}
    before = editable_targets(report_model)
    original_labels = {
        token for text in before.values() for token in citation_tokens(text)
    }
    changed_targets: set[str] = set()
    for operation in plan.operations:
        if operation.operation == "reorder_sections":
            by_id = {
                str(section["id"]): section
                for section in result["synthesis"].get("sections") or []
            }
            result["synthesis"]["sections"] = [by_id[item] for item in operation.order]
            continue
        if operation.operation == "set_format_override":
            overrides[operation.target] = copy.deepcopy(operation.value)
            continue
        replacement = str(operation.replacement or "").strip()
        _validate_replacement_language(replacement, str(result.get("language") or "en"))
        if operation.preserve_citations and Counter(
            citation_tokens(before[operation.target])
        ) != Counter(citation_tokens(replacement)):
            raise RevisionValidationError(
                f"Operation {operation.operation} did not preserve citations in {operation.target}"
            )
        _set_target(result, operation.target, replacement)
        changed_targets.add(operation.target)
    after = editable_targets(result)
    new_labels = {token for text in after.values() for token in citation_tokens(text)}
    if new_labels != original_labels:
        raise RevisionValidationError("Revision changed the report's source-label set")
    before_numbers = Counter(
        token for text in before.values() for token in _LOCKED_NUMBER.findall(text)
    )
    after_numbers = Counter(
        token for text in after.values() for token in _LOCKED_NUMBER.findall(text)
    )
    if before_numbers != after_numbers:
        raise RevisionValidationError("Revision changed a locked numerical finding")
    before_quotes = Counter(
        token for text in before.values() for token in _LOCKED_QUOTE.findall(text)
    )
    after_quotes = Counter(
        token for text in after.values() for token in _LOCKED_QUOTE.findall(text)
    )
    if before_quotes != after_quotes:
        raise RevisionValidationError("Revision changed a locked quotation")
    if not changed_targets and not overrides and not any(
        operation.operation == "reorder_sections" for operation in plan.operations
    ):
        raise RevisionValidationError("Revision plan produced no change")
    return result, overrides


def _reportable(claim: Any) -> bool:
    relevance = float((claim.audit or {}).get("question_relevance", 0.0))
    supporting = int((claim.audit or {}).get("supporting_evidence", 0))
    return claim.status in {"supported", "qualified"} and relevance >= 0.20 and supporting > 0


def _citation_payload(row: Any) -> dict[str, Any]:
    drop_reason = getattr(row, "drop_reason", None)
    return {
        "source_id": str(row.source_id),
        "label": str(row.label),
        "number": int(row.number),
        "cited_sections": list(row.cited_sections or []),
        "offered_sections": list(row.offered_sections or []),
        "claim_ids": list(row.claim_ids or []),
        "evidence_ids": list(row.evidence_ids or []),
        "citation_count": int(row.citation_count),
        "in_bibliography": bool(row.in_bibliography),
        "drop_reason": str(drop_reason or "cited"),
    }


def _zip_with_replacements(original: bytes, replacements: dict[str, bytes]) -> bytes:
    output = io.BytesIO()
    seen: set[str] = set()
    with zipfile.ZipFile(io.BytesIO(original), "r") as source, zipfile.ZipFile(
        output, "w", compression=zipfile.ZIP_DEFLATED
    ) as target:
        for info in source.infolist():
            data = replacements.get(info.filename, source.read(info.filename))
            target.writestr(info, data)
            seen.add(info.filename)
        for name, data in replacements.items():
            if name not in seen:
                target.writestr(name, data)
    return output.getvalue()


def _walk_pptx_shapes(shapes):
    for shape in shapes:
        yield shape
        if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
            yield from _walk_pptx_shapes(shape.shapes)


def office_format_context(data: bytes, media_type: str) -> dict[str, Any]:
    """Describe safe formatting targets without exposing raw OOXML to the planner."""
    if media_type == PPTX_MEDIA_TYPE:
        presentation = Presentation(io.BytesIO(data))
        slides = []
        emu_per_cm = float(PptxCm(1))
        for slide_number, slide in enumerate(presentation.slides, 1):
            shapes = []
            for shape in _walk_pptx_shapes(slide.shapes):
                shape_type = getattr(shape.shape_type, "name", str(shape.shape_type))
                if getattr(shape, "has_chart", False):
                    shape_type = "CHART"
                elif getattr(shape, "has_table", False):
                    shape_type = "TABLE"
                text_preview = ""
                if getattr(shape, "has_text_frame", False):
                    text_preview = " ".join(str(shape.text or "").split())[:160]
                shapes.append(
                    {
                        "name": str(shape.name),
                        "type": str(shape_type),
                        "text": text_preview,
                        "bounds_cm": {
                            "left": round(float(shape.left) / emu_per_cm, 2),
                            "top": round(float(shape.top) / emu_per_cm, 2),
                            "width": round(float(shape.width) / emu_per_cm, 2),
                            "height": round(float(shape.height) / emu_per_cm, 2),
                        },
                    }
                )
            slides.append({"number": slide_number, "shapes": shapes})
        return {
            "pptx": {
                "slide_width_cm": round(
                    float(presentation.slide_width) / emu_per_cm, 2
                ),
                "slide_height_cm": round(
                    float(presentation.slide_height) / emu_per_cm, 2
                ),
                "slides": slides,
            }
        }
    if media_type == DOCX_MEDIA_TYPE:
        document = Document(io.BytesIO(data))
        style_names = {
            paragraph.style.name
            for paragraph in document.paragraphs
            if paragraph.style is not None
        }
        for table in document.tables:
            for row in table.rows:
                for cell in row.cells:
                    style_names.update(
                        paragraph.style.name
                        for paragraph in cell.paragraphs
                        if paragraph.style is not None
                    )
        return {"docx": {"styles": sorted(style_names)}}
    return {}


def _apply_format_overrides(
    docx_data: bytes, pptx_data: bytes, overrides: dict[str, Any]
) -> tuple[bytes, bytes]:
    """Apply the small, validated format grammar without exposing OOXML to the model."""
    docx_items = {
        target: value for target, value in overrides.items() if target.startswith("docx.")
    }
    pptx_items = {
        target: value for target, value in overrides.items() if target.startswith("pptx.")
    }
    if docx_items:
        document = Document(io.BytesIO(docx_data))
        for target, value in docx_items.items():
            _validate_format_override(target, value)
            match = _DOCX_FONT_TARGET.fullmatch(target)
            if match is None:
                raise RevisionValidationError(f"Unsupported DOCX override: {target}")
            style_name = match.group(1)
            try:
                style = document.styles[style_name]
            except KeyError as exc:
                raise RevisionValidationError(f"DOCX style not found: {style_name}") from exc
            style.font.size = DocxPt(float(value))
        output = io.BytesIO()
        document.save(output)
        docx_data = output.getvalue()
    if pptx_items:
        presentation = Presentation(io.BytesIO(pptx_data))
        for target, value in pptx_items.items():
            _validate_format_override(target, value)
            match = _PPTX_BOUNDS_TARGET.fullmatch(target) or _PPTX_FONT_TARGET.fullmatch(target)
            if match is None:
                raise RevisionValidationError(f"Unsupported PPTX override: {target}")
            slide_number = int(match.group(1))
            if slide_number > len(presentation.slides):
                raise RevisionValidationError(f"PPTX slide not found: {slide_number}")
            shape_name = match.group(2)
            shape = next(
                (
                    item
                    for item in _walk_pptx_shapes(
                        presentation.slides[slide_number - 1].shapes
                    )
                    if item.name == shape_name
                ),
                None,
            )
            if shape is None:
                raise RevisionValidationError(
                    f"PPTX shape not found on slide {slide_number}: {shape_name}"
                )
            if _PPTX_BOUNDS_TARGET.fullmatch(target):
                for key, raw in value.items():
                    setattr(shape, key.removesuffix("_cm"), PptxCm(float(raw)))
                if (
                    shape.left < 0
                    or shape.top < 0
                    or shape.width <= 0
                    or shape.height <= 0
                    or shape.left + shape.width > presentation.slide_width
                    or shape.top + shape.height > presentation.slide_height
                ):
                    raise RevisionValidationError(
                        f"PPTX shape exceeds slide bounds: {shape_name}"
                    )
                continue
            if not getattr(shape, "has_text_frame", False):
                raise RevisionValidationError(f"PPTX shape has no text: {shape_name}")
            for paragraph in shape.text_frame.paragraphs:
                for run in paragraph.runs:
                    run.font.size = PptxPt(float(value))
        output = io.BytesIO()
        presentation.save(output)
        pptx_data = output.getvalue()
    return docx_data, pptx_data


def _validate_outputs(outputs: dict[str, tuple[str, bytes]]) -> dict[str, Any]:
    checks: dict[str, Any] = {"files": {}}
    for name, (media_type, data) in outputs.items():
        valid_zip = zipfile.is_zipfile(io.BytesIO(data))
        if not valid_zip:
            raise RevisionValidationError(f"{name} is not a valid ZIP/OOXML container")
        if media_type == DOCX_MEDIA_TYPE:
            document = Document(io.BytesIO(data))
            detail = {"paragraphs": len(document.paragraphs), "tables": len(document.tables)}
        elif media_type == PPTX_MEDIA_TYPE:
            presentation = Presentation(io.BytesIO(data))
            detail = {"slides": len(presentation.slides)}
        else:
            with zipfile.ZipFile(io.BytesIO(data)) as archive:
                detail = {"entries": len(archive.infolist())}
        checks["files"][name] = {"zip": True, **detail}
    checks["passed"] = True
    return checks


class DocumentRevisionService:
    def __init__(
        self,
        repo: Repository,
        store: ObjectStore,
        *,
        llm: LLMProvider | None = None,
        settings: Any | None = None,
    ) -> None:
        self.repo = repo
        self.store = store
        self.llm = llm
        self.settings = settings

    async def ensure_initial_revision(self, run_id: str):
        current = await self.repo.current_document_revision(run_id)
        if current is not None:
            return current
        run = await self.repo.get_run(run_id)
        if run is None or run.status not in TERMINAL_RUN_STATUSES:
            raise RevisionConflict("Only completed research runs can be revised")
        artifacts = await self.repo.list_artifacts(run_id)
        by_name = {artifact.name: artifact for artifact in artifacts}
        manifest_artifact = by_name.get("10_reproducibility_manifest.json")
        if manifest_artifact is None:
            raise RevisionConflict("The run has no reproducibility manifest")
        manifest = json.loads((await self.store.get(manifest_artifact.object_key)).decode("utf-8"))
        version_rows = []
        for artifact in artifacts:
            if artifact.media_type not in {DOCX_MEDIA_TYPE, PPTX_MEDIA_TYPE} and (
                artifact.name not in VERSIONED_BUNDLES
            ):
                continue
            data = await self.store.get(artifact.object_key)
            version_rows.append(
                {
                    "logical_name": artifact.name,
                    "media_type": artifact.media_type,
                    "object_key": artifact.object_key,
                    "sha256": hashlib.sha256(data).hexdigest(),
                    "size_bytes": len(data),
                }
            )
        if not any(item["media_type"] == DOCX_MEDIA_TYPE for item in version_rows) or not any(
            item["media_type"] == PPTX_MEDIA_TYPE for item in version_rows
        ):
            raise RevisionConflict("The run must have both DOCX and PPTX artifacts")
        citations = [
            _citation_payload(row) for row in await self.repo.list_report_citations(run_id)
        ]
        return await self.repo.bootstrap_document_revision(
            run_id,
            report_model=report_model_from_manifest(manifest),
            artifact_versions=version_rows,
            citations=citations,
            requested_by=str(run.owner_id or "system"),
        )

    async def plan(self, revision_id: str):
        if self.llm is None:
            raise RuntimeError("Revision planning requires an LLM provider")
        revision = await self.repo.get_document_revision_by_id(revision_id, lock=True)
        if revision.status == RevisionStatus.CANCELLED.value:
            return revision
        if revision.status != RevisionStatus.QUEUED.value or revision.edit_plan:
            raise RevisionConflict(f"Cannot plan revision in {revision.status}")
        parent = await self.repo.get_document_revision(
            revision.run_id, str(revision.parent_revision_id)
        )
        await self.repo.update_document_revision(
            revision.run_id,
            revision.id,
            status=RevisionStatus.PLANNING.value,
            error=None,
        )
        format_targets: dict[str, Any] = {}
        if _FORMAT_FEEDBACK.search(revision.feedback):
            parent_versions = await self.repo.list_artifact_versions(
                revision.run_id, revision_id=parent.id
            )
            target_version = next(
                (
                    item
                    for item in parent_versions
                    if item.logical_name == revision.target_artifact_name
                ),
                None,
            )
            if target_version is not None and target_version.media_type in {
                DOCX_MEDIA_TYPE,
                PPTX_MEDIA_TYPE,
            }:
                format_targets = office_format_context(
                    await self.store.get(target_version.object_key),
                    target_version.media_type,
                )
        plan, metrics = await generate_revision_plan(
            self.llm,
            parent.report_model,
            revision.feedback,
            format_targets=format_targets,
        )
        current = await self.repo.get_document_revision(
            revision.run_id, revision.id, lock=True
        )
        if current.status == RevisionStatus.CANCELLED.value:
            return current
        if current.status != RevisionStatus.PLANNING.value:
            raise RevisionConflict(f"Revision left planning state: {current.status}")
        status = (
            RevisionStatus.CLARIFYING.value
            if plan.requires_clarification
            else RevisionStatus.AWAITING_PLAN_APPROVAL.value
        )
        return await self.repo.update_document_revision(
            revision.run_id,
            revision.id,
            status=status,
            edit_plan=plan.model_dump(mode="json"),
            model_id=str(getattr(self.settings, "llm_model", "") or "unknown"),
            prompt_version=PLAN_PROMPT_VERSION,
            usage={"calls": metrics},
            channel_state={},
        )

    async def _figures(self, run_id: str) -> tuple[list[FigureObservation], list[GeneratedResearchFigure]]:
        artifacts = {item.name: item for item in await self.repo.list_artifacts(run_id)}
        manifest_artifact = artifacts.get("17_figure_observations.json")
        if manifest_artifact is None:
            return [], []
        try:
            payload = json.loads((await self.store.get(manifest_artifact.object_key)).decode("utf-8"))
        except (OSError, ValueError, UnicodeDecodeError):
            return [], []
        observations = []
        observation_fields = {item.name for item in fields(FigureObservation)}
        for item in payload.get("observations") or []:
            observations.append(
                FigureObservation(
                    **{key: value for key, value in item.items() if key in observation_fields}
                )
            )
        generated = []
        figure_fields = {item.name for item in fields(GeneratedResearchFigure)} - {"data"}
        for item in payload.get("generated_figures") or []:
            artifact = artifacts.get(str(item.get("name") or ""))
            if artifact is None:
                continue
            generated.append(
                GeneratedResearchFigure(
                    data=await self.store.get(artifact.object_key),
                    **{key: value for key, value in item.items() if key in figure_fields},
                )
            )
        return observations, generated

    async def render(self, revision_id: str):
        revision = await self.repo.get_document_revision_by_id(revision_id, lock=True)
        if revision.status == RevisionStatus.CANCELLED.value:
            return revision
        if revision.status != RevisionStatus.QUEUED.value or not revision.edit_plan:
            raise RevisionConflict(f"Cannot render revision in {revision.status}")
        plan = RevisionPlan.model_validate(revision.edit_plan)
        parent = await self.repo.get_document_revision(
            revision.run_id, str(revision.parent_revision_id)
        )
        report_model, new_overrides = apply_revision_plan(parent.report_model, plan)
        format_overrides = {**(parent.format_overrides or {}), **new_overrides}
        await self.repo.update_document_revision(
            revision.run_id,
            revision.id,
            status=RevisionStatus.RENDERING.value,
            error=None,
        )
        protocol = ResearchProtocol.model_validate(report_model["protocol"])
        package = synthesis_package_from_model(report_model)
        coverage = CoverageMetrics.model_validate(report_model.get("coverage") or {})
        sources = await self.repo.list_sources(revision.run_id)
        claims = await self.repo.list_claims(revision.run_id)
        evidence = await self.repo.list_evidence(revision.run_id)
        evidence_by_claim: dict[str, list[tuple[Any, Any]]] = {}
        for claim, link, source in evidence:
            evidence_by_claim.setdefault(claim.id, []).append((link, source))
        reportable = sorted(
            [claim for claim in claims if _reportable(claim)],
            key=lambda claim: (
                claim.status == "supported",
                claim.importance == "major",
                float((claim.audit or {}).get("question_relevance", 0.0)),
            ),
            reverse=True,
        )
        observations, generated_figures = await self._figures(revision.run_id)
        formula_displays = {}
        if self.settings is not None and getattr(
            self.settings, "formula_resolution_enabled", False
        ):
            version_ids = sorted({str(link.source_version_id) for _, link, _ in evidence})
            passage_texts = {
                passage.id: passage.text
                for passage in await self.repo.list_passages(revision.run_id, version_ids)
            }
            formula_displays = await load_formula_displays(
                repo=self.repo,
                store=self.store,
                links=[link for _, link, _ in evidence],
                settings=self.settings,
                passage_texts=passage_texts,
            )
        common = {
            "run_id": revision.run_id,
            "title": str(report_model["title"]),
            "question": str(report_model["question"]),
            "language": str(report_model["language"]),
            "coverage": coverage.model_dump(),
            "sources": sources,
            "claims": claims,
            "reportable_claims": reportable,
            "evidence_by_claim": evidence_by_claim,
            "executive_summary": package.executive_summary,
            "narrative": package.narrative,
            "uncertainty": package.uncertainty,
            "scope": protocol.scope.model_dump(mode="json"),
            "sub_questions": protocol.sub_questions_for_report(),
            "connector_ids": protocol.connectors.included_connectors,
            "research_mode": protocol.research_mode,
            "synthesis_package": package,
            "figure_observations": observations,
            "research_figures": generated_figures,
        }
        parent_versions = await self.repo.list_artifact_versions(
            revision.run_id, revision_id=parent.id
        )
        parent_by_type = {item.media_type: item for item in parent_versions}
        docx_parent = parent_by_type.get(DOCX_MEDIA_TYPE)
        pptx_parent = parent_by_type.get(PPTX_MEDIA_TYPE)
        if docx_parent is None or pptx_parent is None:
            raise RevisionConflict("The base revision does not contain both Office formats")
        changed_scopes = {item.format_scope for item in plan.operations}
        render_both = "all" in changed_scopes or not changed_scopes
        outputs: dict[str, tuple[str, bytes]] = {}
        if render_both or "docx" in changed_scopes:
            word = build_word_report(**common, formula_displays=formula_displays)
            docx_data = word.document
            citations = [item.model_dump(mode="json") for item in word.citations]
        else:
            docx_data = await self.store.get(docx_parent.object_key)
            citations = [
                item.model_dump(mode="json")
                for item in [
                    ReportCitation.model_validate(row.payload)
                    for row in await self.repo.list_revision_citations(
                        revision.run_id, parent.id
                    )
                ]
            ]
        if render_both or "pptx" in changed_scopes:
            presentation = build_presentation_report(**common)
            # A re-rendered deck is the builder's again; without this, every revision
            # would silently undo the agent's edits.
            pptx_data = await polish_and_record(
                self.repo,
                revision.run_id,
                presentation.document,
                language=protocol.report_language,
                settings=self.settings or get_settings(),
                stage="revision",
                revision_id=revision.id,
            )
        else:
            pptx_data = await self.store.get(pptx_parent.object_key)
        docx_data, pptx_data = _apply_format_overrides(
            docx_data, pptx_data, format_overrides
        )
        current = await self.repo.get_document_revision(
            revision.run_id, revision.id, lock=True
        )
        if current.status == RevisionStatus.CANCELLED.value:
            return current
        if current.status != RevisionStatus.RENDERING.value:
            raise RevisionConflict(f"Revision left rendering state: {current.status}")
        outputs[docx_parent.logical_name] = (DOCX_MEDIA_TYPE, docx_data)
        outputs[pptx_parent.logical_name] = (PPTX_MEDIA_TYPE, pptx_data)
        replacements = {
            docx_parent.logical_name: docx_data,
            pptx_parent.logical_name: pptx_data,
        }
        parent_by_name = {item.logical_name: item for item in parent_versions}
        for bundle_name in VERSIONED_BUNDLES:
            base_bundle = parent_by_name.get(bundle_name)
            if base_bundle is None:
                continue
            outputs[bundle_name] = (
                "application/zip",
                _zip_with_replacements(
                    await self.store.get(base_bundle.object_key), replacements
                ),
            )
        await self.repo.update_document_revision(
            revision.run_id,
            revision.id,
            status=RevisionStatus.VALIDATING.value,
        )
        validation = _validate_outputs(outputs)
        version_rows = []
        for name, (media_type, data) in outputs.items():
            object_key = f"runs/{revision.run_id}/revisions/{revision.id}/{name}"
            await self.store.put(object_key, data, media_type)
            version_rows.append(
                {
                    "logical_name": name,
                    "parent_version_id": (
                        parent_by_name[name].id if name in parent_by_name else None
                    ),
                    "media_type": media_type,
                    "object_key": object_key,
                    "sha256": hashlib.sha256(data).hexdigest(),
                    "size_bytes": len(data),
                }
            )
        return await self.repo.save_revision_outputs(
            revision.run_id,
            revision.id,
            report_model=report_model,
            format_overrides=format_overrides,
            validation=validation,
            artifact_versions=version_rows,
            citations=citations,
        )
