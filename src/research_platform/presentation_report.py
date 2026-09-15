"""Deterministic PowerPoint (PPTX) report renderer for completed research runs.

Draws the Cansağlığı research report deck ("Editoryal Ray" design) from audited run data:
synthesis, figures, claims, citations and the literature landscape.

Every slide is drawn here from fixed geometry, and every text box is measured with the
deck's font before it is placed (`presentation_layout`). Nothing relies on PowerPoint's
shrink-on-overflow, which never runs for a generated file. The rules the layout keeps:

- one font size per role; no text is shrunk to fit and none is cut;
- prose that does not fit continues on a next slide, split between sentences;
- a heading longer than three lines steps down one size, and the body below it moves down;
- appendix tables show the rows that fit one slide and name where the full list is.
"""

from __future__ import annotations

import io
import re
from collections import Counter
from dataclasses import dataclass, field, replace
from datetime import UTC, datetime
from pathlib import Path
from typing import Any
from urllib.parse import urlparse

import pptx
from PIL import Image
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LABEL_POSITION
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Pt

from .figure_analysis import FigureObservation, GeneratedResearchFigure
from .presentation_layout import (
    FONT_NAME,
    Flow,
    Placed,
    TextStyle,
    fit_lines,
    fit_sentences,
    line_count,
    paginate,
    sentences,
    text_height,
    text_width,
)
from .report_synthesis import (
    _REPORT_WITHOUT_SUMMARY,
    SynthesisPackage,
    citation_tokens,
    cited_labels,
)
from .schemas import ReportCitation
from .word_report import _collect_citations, _format_date

REPORT_PIPELINE_VERSION = "0.24.0"
PRESENTATION_REPORT_FALLBACK = "16_research_report.pptx"
_SAFE_LABEL = re.compile(r"[A-Za-z0-9_]{1,64}")

TEMPLATES = Path(__file__).parent / "templates"
TEMPLATE_PATH = TEMPLATES / "Cansagligi_Arastirma_Raporu_Sablonu_v2.pptx"
LOGO_WHITE = TEMPLATES / "brand" / "logo-beyaz.png"
LOGO_COLOR = TEMPLATES / "brand" / "logo-renkli.png"

# Colours of the design canvas ("Cansağlığı Rapor Sunumu", direction B).
RED = "D90919"
INK = "242424"
GREY = "555555"
MUTED = "646464"
RULE = "E4E2E0"
RAIL = "F4F2F0"
WHITE = "FFFFFF"

# Geometry in points on the 960 x 540 pt (13.33 x 7.5 in) page.
SLIDE_W, SLIDE_H = 960.0, 540.0
RAIL_W = 225.0
RAIL_X = 30.0
RAIL_INNER_W = RAIL_W - 2 * RAIL_X
CONTENT_X = 273.0
CONTENT_W = 627.0
TOP = 48.0
FOOTER_TOP = 486.0
BODY_BOTTOM = 474.0
HEADING_GAP = 21.0
COLUMN_GAP = 30.0
COVER_PANEL_W = 336.0

HEADING = TextStyle(24, bold=True, leading=1.25)
HEADING_LONG = TextStyle(20, bold=True, leading=1.25)
HEADING_SMALL = TextStyle(16, bold=True, leading=1.25)
BODY = TextStyle(16, leading=1.5)
LEAD = TextStyle(18, leading=1.5)
LIST = TextStyle(14, leading=1.35)
SMALL = TextStyle(12, color=MUTED, leading=1.3)
SMALL_INK = TextStyle(12, leading=1.3)
SMALL_BOLD = TextStyle(12, bold=True, leading=1.3)
EYEBROW = TextStyle(12, bold=True, color=RED, leading=1.2, tracking=1.4)
RAIL_NUMBER = TextStyle(60, bold=True, color=RED, leading=1.0)
RAIL_LABEL = TextStyle(14, color=GREY, leading=1.35)
COVER_TITLE = TextStyle(37.5, bold=True, leading=1.1)
COVER_TITLE_LONG = TextStyle(30, bold=True, leading=1.1)
# Whole sizes the cover title steps through, each with the most lines it may take there.
_COVER_TITLE_STEPS = (
    (COVER_TITLE, 4),
    (COVER_TITLE_LONG, 6),
    (TextStyle(24, bold=True, leading=1.15), 99),
    (TextStyle(20, bold=True, leading=1.15), 99),
    (TextStyle(16, bold=True, leading=1.2), 99),
    (TextStyle(12, bold=True, leading=1.25), 99),
)


def presentation_report_name(label: str | None) -> str:
    """The PowerPoint report's file name, derived from the run's topic handle."""
    from .scoping import LABEL_MAX_LENGTH, slugify

    handle = (label or "").strip()
    if not _SAFE_LABEL.fullmatch(handle):
        handle = slugify(handle, max_length=LABEL_MAX_LENGTH)
    return f"16_{handle}_report.pptx" if handle else PRESENTATION_REPORT_FALLBACK


@dataclass(frozen=True)
class PresentationReportResult:
    document: bytes
    figures: dict[str, bytes] = field(default_factory=dict)
    citations: list[ReportCitation] = field(default_factory=list)


_TEXT = {
    "report": ("ARAŞTIRMA RAPORU", "RESEARCH REPORT"),
    "report_date": ("RAPOR TARİHİ", "REPORT DATE"),
    "run_id": ("Çalışma kimliği", "Run ID"),
    "toc": ("İçindekiler", "Contents"),
    "toc_report": ("RAPOR", "REPORT"),
    "toc_appendix": ("EKLER", "APPENDICES"),
    "summary": ("Özet", "Summary"),
    "frame": ("Araştırma çerçevesi", "Research frame"),
    "question_scope": ("Ana soru ve kapsam", "Main question and scope"),
    "question": ("ARAŞTIRMA SORUSU", "RESEARCH QUESTION"),
    "date_range": ("Tarih aralığı", "Date range"),
    "mode": ("Araştırma modu", "Research mode"),
    "intent": ("Kapsam gerekçesi", "Scope rationale"),
    "sources_row": ("Kaynaklar", "Sources"),
    "claims_row": ("Denetlenen iddialar", "Audited claims"),
    "sub_questions": ("Alt sorular ve kapsam sınırları", "Sub-questions and scope limits"),
    "near": ("YAKIN AMA KAPSAM DIŞI", "NEAR BUT OUT OF SCOPE"),
    "themes": ("Tematik kanıt sentezi", "Thematic evidence synthesis"),
    "assessment_section": ("Değerlendirme ve sonuç", "Assessment and conclusion"),
    "toc_assessment": (
        "Çalışmalar arası değerlendirme ve sonuç",
        "Cross-study assessment and conclusion",
    ),
    "assessment": ("Çalışmalar arası değerlendirme", "Cross-study assessment"),
    "conclusion": ("Sonuç", "Conclusion"),
    "uncertainty": ("Belirsizlikler ve araştırma boşlukları", "Uncertainties and research gaps"),
    "consensus": ("ORTAK YÖN", "CONSENSUS"),
    "disagreements": ("AYRIŞMA", "DISAGREEMENT"),
    "implications": ("ARAŞTIRMA AÇISINDAN ANLAMI", "WHAT IT MEANS FOR THE RESEARCH"),
    "figures_section": ("Kaynak figürleri", "Source figures"),
    "figure": ("FİGÜR", "FIGURE"),
    "limitation": ("Sınır", "Limitation"),
    "source": ("Kaynak", "Source"),
    "rights": ("Telif", "Rights"),
    "appendix_a": ("Ek A", "A"),
    "appendix_b": ("Ek B", "B"),
    "appendix_c": ("Ek C", "C"),
    "appendix_d": ("Ek D", "D"),
    "method_section": ("Yöntem ve kapsam", "Method and scope"),
    "method": ("Yöntem, kapsam ve yeniden üretilebilirlik", "Method, scope and reproducibility"),
    "landscape": ("Literatürün konu haritası", "Literature topic map"),
    "catalog": ("Kaynak kataloğu", "Source catalog"),
    "claims": ("Denetlenmiş iddia kaydı", "Audited claim register"),
    "metric": ("Ölçüt", "Metric"),
    "value": ("Değer", "Value"),
    "continued": (" (devam)", " (continued)"),
    "sources_label": ("Kaynaklar", "Sources"),
    "more_sources": ("kaynak daha", "more sources"),
    "thanks": ("Teşekkürler", "Thank you"),
    "foundation": ("Canan Bayraktar Toplum Sağlığı Vakfı", "Canan Bayraktar Community Health Foundation"),
    "no_consensus": ("Bu tema için ortak bir yön bildirilmedi.", "No consensus reported."),
    "no_disagreement": (
        "Kaynaklar arasında doğrudan bir çelişki bildirilmedi.",
        "No direct contradiction reported.",
    ),
    "no_implication": (
        "Bu tema için özel bir araştırma çıkarımı belirtilmedi.",
        "No specific research implication reported.",
    ),
    "no_assessment": (
        "Çalışmalar arası değerlendirme bu raporda ayrıca yer almıyor.",
        "No separate cross-study assessment is included in this report.",
    ),
    "no_conclusion": (
        "Sonuç bölümü bu raporda ayrıca yer almıyor.",
        "No separate conclusion is included in this report.",
    ),
    "no_uncertainty": (
        "Belirsizlik bölümü bu raporda ayrıca yer almıyor.",
        "No separate uncertainty section is included in this report.",
    ),
    "no_sub_questions": (
        "Ana araştırma sorusu bütüncül olarak ele alınmıştır.",
        "The primary research question was addressed holistically.",
    ),
    "no_near_scope": (
        "Kapsam dışı bırakılan özel bir sınır çalışma bildirilmemiştir.",
        "No specific near-scope excluded studies were flagged.",
    ),
    "no_chart": (
        "Araştırma katkısı dağılımı bu koşuda üretilmedi.",
        "Research contribution distribution not available.",
    ),
    "no_map": ("Tema-kanıt haritası bu koşuda üretilmedi.", "Theme-evidence map not available."),
    "no_sources": ("Araştırmada kayıtlı kaynak bulunamadı.", "No retained sources found."),
    "no_claims": ("Denetlenmiş iddia kaydı bulunamadı.", "No audited claims found."),
}

_METHOD_STEPS = {
    True: [
        "Keşif: Çoklu akademik ve web arama bağlayıcıları üzerinden literatür tarandı.",
        "Edinim: Tam metinler ve açık erişimli içerikler tekilleştirilerek edinildi.",
        "Normalizasyon: URL, içerik özeti ve üst veriler kaydedildi.",
        "Kanıt: Pasajlar atomik iddialarla eşleştirildi ve güven puanlandı.",
        "Sentez: Yalnızca denetim kapısını geçen iddialar rapora alındı.",
    ],
    False: [
        "Discovery: Federated search across academic and web connectors.",
        "Acquisition: Open access content deduplicated and indexed.",
        "Normalization: Content hashing and metadata validation.",
        "Evidence: Claim linking with entailment scoring.",
        "Synthesis: Only audited findings entered the report.",
    ],
}

_MODES = {
    "literature_scan": ("Literatür taraması", "Literature scan"),
    "focused_answer": ("Odaklı yanıt", "Focused answer"),
}
_FAMILIES = {
    "web": ("web", "web"),
    "academic": ("akademik", "academic"),
    "code_data": ("kod ve veri", "code and data"),
    "news": ("haber", "news"),
}
_FAMILY_SHORT = {
    "web": ("Web", "Web"),
    "academic": ("Akademik", "Academic"),
    "code_data": ("Kod/veri", "Code/data"),
    "news": ("Haber", "News"),
}
_SCOPE_ROLES = {
    "near_scope": ("Yakın kapsam", "Near scope"),
    "excluded": ("Kapsam dışı", "Excluded"),
}
_STATUSES = {
    "supported": ("Destekleniyor", "Supported"),
    "qualified": ("Sınırlı destek", "Qualified"),
    "contested": ("Tartışmalı", "Contested"),
}
# Site names search engines append to page titles; the catalog shows the domain separately.
_TITLE_SUFFIX = re.compile(
    r"(?:\s+[|·]\s+.*|\s+-\s+(?:PMC|arXiv|ScienceDirect|PubMed|ResearchGate|MDPI|GitHub|ejrai))$",
    re.IGNORECASE,
)
_FIGURE_NUMBER = re.compile(r"^\s*(?:Şekil|Sekil|Figure|Figür|Fig\.?)\s*\d+[a-z]?\b\s*(?:[:.\-–—]\s*|$)", re.IGNORECASE)


def _is_turkish(language: str) -> bool:
    return language.lower().strip().startswith("tr")


def _clean(value: Any) -> str:
    return " ".join(str(value or "").replace("\x00", "").split())


def _label_number(label: str) -> int:
    digits = label.lstrip("[S").rstrip("]")
    return int(digits) if digits.isdigit() else 10**6


def _cited(*texts: str) -> list[str]:
    labels = {token.strip("[]") for text in texts for token in citation_tokens(text or "")}
    return sorted(labels, key=_label_number)


def _percent(value: Any, turkish: bool) -> str:
    if not isinstance(value, (int, float)) or isinstance(value, bool):
        return "—"
    share = float(value) * 100
    number = f"{share:.0f}" if share >= 10 or share == round(share) else f"{share:.1f}"
    if turkish:
        return f"%{number.replace('.', ',')}"
    return f"{number}%"


def _decimal(value: Any, turkish: bool) -> str:
    try:
        rendered = f"{float(value):.2f}"
    except (TypeError, ValueError):
        return "—"
    return rendered.replace(".", ",") if turkish else rendered


def _source_evidence_counts(
    sources: list[Any], evidence_by_claim: dict[str, list[tuple[Any, Any]]]
) -> dict[str, int]:
    """How many claims each source supports."""
    counts: dict[str, int] = {str(s.id): 0 for s in sources}
    for links in evidence_by_claim.values():
        seen: set[str] = set()
        for _, source in links:
            source_id = str(source.id)
            if source_id not in seen:
                counts[source_id] = counts.get(source_id, 0) + 1
                seen.add(source_id)
    return counts


def _claim_sources(
    claim_id: str,
    evidence_by_claim: dict[str, list[tuple[Any, Any]]],
    source_numbers: dict[Any, int],
) -> list[int]:
    numbers: list[int] = []
    for _, source in evidence_by_claim.get(str(claim_id), []):
        number = source_numbers.get(source.id)
        if number is not None and number not in numbers:
            numbers.append(number)
    return numbers


def _section_text(section: Any) -> list[str]:
    """A theme's paragraphs, headed by the reader note when it carries one."""
    paragraphs = []
    note = _clean(getattr(section, "reader_note", ""))
    if note:
        paragraphs.append(note)
    paragraphs.extend(_paragraphs(getattr(section, "synthesis", "")))
    return paragraphs


def _paragraphs(text: Any) -> list[str]:
    return [_clean(part) for part in re.split(r"\n\s*\n", str(text or "")) if _clean(part)]


_HEADING_STEPS = ((HEADING, 3), (HEADING_LONG, 4), (HEADING_SMALL, None))
# Room a heading must leave for its body. A heading that would leave less is laid out as the
# body's first element instead, so it can continue onto the next slide like any text.
_MIN_BODY_H = 4 * BODY.line_height
_FIGURE_IMAGE_MIN_H = 150.0
_FIGURE_TITLE_STEPS = (
    (HEADING, 3),
    (HEADING_LONG, 4),
    (HEADING_SMALL, 6),
    (TextStyle(12, bold=True, leading=1.25), None),
)


def _heading_style(text: str, width: float = CONTENT_W) -> TextStyle:
    for style, max_lines in _HEADING_STEPS:
        if max_lines is None or line_count(text, style, width) <= max_lines:
            return style
    return HEADING_SMALL


def _heading_fits(heading: str) -> bool:
    height = text_height(heading, _heading_style(heading), CONTENT_W)
    return TOP + height + HEADING_GAP + _MIN_BODY_H <= BODY_BOTTOM


def _heading_flows(heading: str) -> list[Flow]:
    """A heading too tall to sit above its body, laid out as the body's first element."""
    return [Flow(heading, HEADING_SMALL, CONTENT_W, pad_bottom=HEADING_GAP, name="content_heading")]


def _image_size(data: bytes) -> tuple[int, int] | None:
    try:
        with Image.open(io.BytesIO(data)) as image:
            return image.size
    except (OSError, ValueError, TypeError):
        return None


@dataclass
class _Page:
    """One slide to draw, decided before any slide exists so the contents know page numbers."""

    draw: Any
    toc_key: str = ""


class _Deck:
    def __init__(self, template: Path, turkish: bool):
        self.prs = pptx.Presentation(str(template))
        self.layout = self.prs.slide_layouts[0]
        self.turkish = turkish
        self.logo_white = LOGO_WHITE.read_bytes()
        self.logo_color = LOGO_COLOR.read_bytes()

    def t(self, key: str) -> str:
        return _TEXT[key][0 if self.turkish else 1]

    # -- primitives -------------------------------------------------------------------------

    def slide(self) -> Any:
        slide = self.prs.slides.add_slide(self.layout)
        for shape in list(slide.placeholders):
            shape._element.getparent().remove(shape._element)
        return slide

    def text(
        self,
        slide: Any,
        name: str,
        x: float,
        y: float,
        width: float,
        height: float,
        text: str,
        style: TextStyle,
        align: Any = None,
        anchor: Any = MSO_ANCHOR.TOP,
    ) -> Any:
        shape = slide.shapes.add_textbox(Pt(x), Pt(y), Pt(width), Pt(max(height, style.line_height)))
        shape.name = name
        frame = shape.text_frame
        frame.word_wrap = True
        frame.auto_size = MSO_AUTO_SIZE.NONE
        frame.vertical_anchor = anchor
        body = frame._txBody.find(qn("a:bodyPr"))
        for inset in ("lIns", "tIns", "rIns", "bIns"):
            body.set(inset, "0")
        paragraph = frame.paragraphs[0]
        if align is not None:
            paragraph.alignment = align
        paragraph.line_spacing = Pt(style.line_height)
        paragraph.space_before = Pt(0)
        paragraph.space_after = Pt(0)
        run = paragraph.add_run()
        run.text = text
        font = run.font
        font.name = FONT_NAME
        font.size = Pt(style.size)
        font.bold = style.bold
        font.color.rgb = RGBColor.from_string(style.color)
        properties = run._r.get_or_add_rPr()
        properties.set("lang", "tr-TR" if self.turkish else "en-US")
        if style.tracking:
            properties.set("spc", str(round(style.tracking * 100)))
        return shape

    def rect(self, slide: Any, name: str, x: float, y: float, width: float, height: float, color: str) -> Any:
        shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Pt(x), Pt(y), Pt(width), Pt(height))
        shape.name = name
        shape.fill.solid()
        shape.fill.fore_color.rgb = RGBColor.from_string(color)
        shape.line.fill.background()
        style = shape._element.find(qn("p:style"))
        if style is not None:
            shape._element.remove(style)
        return shape

    def picture(
        self, slide: Any, name: str, data: bytes, x: float, y: float, width: float, height: float
    ) -> Any | None:
        size = _image_size(data)
        if size is None:
            return None
        image_w, image_h = size
        scale = min(width / max(1, image_w), height / max(1, image_h))
        draw_w, draw_h = image_w * scale, image_h * scale
        shape = slide.shapes.add_picture(
            io.BytesIO(data),
            Pt(x + (width - draw_w) / 2),
            Pt(y + (height - draw_h) / 2),
            Pt(draw_w),
            Pt(draw_h),
        )
        shape.name = name
        return shape

    # -- shared slide parts -----------------------------------------------------------------

    def rail(self, slide: Any, number: str, label: str) -> None:
        self.rect(slide, "rail", 0, 0, RAIL_W, SLIDE_H, RAIL)
        number_style = RAIL_NUMBER if text_width(number, RAIL_NUMBER) <= RAIL_INNER_W else HEADING
        self.text(slide, "rail_number", RAIL_X, TOP, RAIL_INNER_W, number_style.line_height, number, number_style)
        label_y = TOP + number_style.line_height + 9
        self.text(
            slide,
            "rail_label",
            RAIL_X,
            label_y,
            RAIL_INNER_W,
            text_height(label, RAIL_LABEL, RAIL_INNER_W),
            label,
            RAIL_LABEL,
        )
        self.logo(slide)

    def logo(self, slide: Any) -> None:
        width = 150.0
        size = _image_size(self.logo_color) or (514, 109)
        height = width * size[1] / size[0]
        self.picture(slide, "logo", self.logo_color, RAIL_X, SLIDE_H - 36 - height, width, height)

    def footer(self, slide: Any, page: int, left: str = "") -> None:
        self.rect(slide, "footer_rule", CONTENT_X, FOOTER_TOP, CONTENT_W, 0.75, RULE)
        number_w = 40.0
        if left:
            self.text(
                slide,
                "footer_note",
                CONTENT_X,
                FOOTER_TOP + 6,
                CONTENT_W - number_w - 12,
                42,
                left,
                SMALL,
                anchor=MSO_ANCHOR.MIDDLE,
            )
        self.text(
            slide,
            "slide_number",
            CONTENT_X + CONTENT_W - number_w,
            FOOTER_TOP + 6,
            number_w,
            42,
            str(page),
            SMALL,
            align=PP_ALIGN.RIGHT,
            anchor=MSO_ANCHOR.MIDDLE,
        )

    def citation_note(self, labels: list[str]) -> str:
        if not labels:
            return ""
        prefix = f"{self.t('sources_label')}  "
        width = CONTENT_W - 52
        note = prefix + " · ".join(labels)
        if line_count(note, SMALL, width) <= 2:
            return note
        for keep in range(len(labels) - 1, 0, -1):
            note = f"{prefix}{' · '.join(labels[:keep])} · +{len(labels) - keep} {self.t('more_sources')}"
            if line_count(note, SMALL, width) <= 2:
                return note
        return f"{prefix}+{len(labels)} {self.t('more_sources')}"

    def heading(self, slide: Any, text: str, y: float = TOP, width: float = CONTENT_W) -> float:
        style = _heading_style(text, width)
        height = text_height(text, style, width)
        self.text(slide, "content_heading", CONTENT_X, y, width, height, text, style)
        return y + height

    def draw_flows(self, slide: Any, pieces: list[Placed], x: float, top: float) -> None:
        for index, piece in enumerate(pieces, 1):
            flow = piece.flow
            y = top + piece.top
            name = flow.name if index == 1 else f"{flow.name}_{index}"
            if flow.strong_rule or flow.rule:
                self.rect(
                    slide,
                    f"{name}_rule",
                    x,
                    y,
                    flow.width,
                    1.5 if flow.strong_rule else 0.75,
                    INK if flow.strong_rule else RULE,
                )
            y += flow.pad_top
            if flow.label and flow.label_style is not None and not piece.continued:
                label_h = text_height(flow.label, flow.label_style, flow.width)
                self.text(slide, f"{name}_label", x, y, flow.width, label_h, flow.label, flow.label_style)
                y += label_h + flow.label_gap
            if flow.marker and not piece.continued:
                marker_style = replace(flow.style, bold=True, color=RED)
                self.text(slide, f"{name}_marker", x, y, flow.marker_width, flow.style.line_height, flow.marker, marker_style)
            if piece.text:
                height = text_height(piece.text, flow.style, flow.text_width)
                self.text(slide, name, x + flow.marker_width, y, flow.text_width, height, piece.text, flow.style)

    def save(self) -> bytes:
        output = io.BytesIO()
        self.prs.save(output)
        return output.getvalue()


def _body_height(heading: str | None) -> tuple[float, float]:
    """Top and height of the body area under a heading."""
    if not heading:
        return TOP, BODY_BOTTOM - TOP
    top = TOP + text_height(heading, _heading_style(heading), CONTENT_W) + HEADING_GAP
    return top, BODY_BOTTOM - top


def build_presentation_report(
    *,
    run_id: str,
    title: str,
    question: str,
    language: str,
    coverage: dict[str, Any],
    sources: list[Any],
    claims: list[Any],
    reportable_claims: list[Any],
    evidence_by_claim: dict[str, list[tuple[Any, Any]]],
    executive_summary: str,
    narrative: str,
    uncertainty: str,
    scope: dict[str, Any] | None = None,
    sub_questions: list[str] | None = None,
    connector_ids: list[str] | None = None,
    research_mode: str = "literature_scan",
    synthesis_package: SynthesisPackage | None = None,
    figure_observations: list[FigureObservation] | None = None,
    research_figures: list[GeneratedResearchFigure] | None = None,
    template_path: Path | str | None = None,
    figures: dict[str, bytes] | None = None,
) -> PresentationReportResult:
    """Build the branded PowerPoint research report for a completed run."""
    turkish = _is_turkish(language)
    template = Path(template_path) if template_path is not None else TEMPLATE_PATH
    if not template.exists():
        raise FileNotFoundError(f"PowerPoint report template not found: {template}")
    deck = _Deck(template, turkish)
    t = deck.t
    package = synthesis_package
    compact = package is not None and getattr(package, "report_mode", "") == "compact"
    sections = list(getattr(package, "sections", [])) if package is not None and not compact else []
    source_numbers = {source.id: index for index, source in enumerate(sources, 1)}
    citations = list(
        _collect_citations(
            sources=sources,
            source_numbers=source_numbers,
            evidence_by_claim=evidence_by_claim,
            reportable_claims=reportable_claims,
            package=package,
            turkish=turkish,
        )
    )
    figures_map = dict(figures or {})

    pages: list[_Page] = []

    # Cover, then the contents. The contents are drawn last, once page numbers are known;
    # how many slides they take depends only on their text, so it is decided here.
    pages.append(_Page(lambda slide, page: _draw_cover(deck, slide, title, run_id)))
    contents_style, contents_chunks = _plan_contents(_contents_rows(deck, compact, sections))
    toc_index = len(pages)
    pages.extend(_Page(None) for _ in contents_chunks)

    # Section 1: summary
    summary = (getattr(package, "executive_summary", "") if package is not None else "") or executive_summary
    if not _clean(summary):
        # The synthesis layer's own wording, so every surface says the same plain thing.
        summary = _REPORT_WITHOUT_SUMMARY[turkish]
    _add_prose(pages, deck, "1", t("summary"), None, _paragraphs(summary), LEAD, "summary", "summary")

    # Section 2: research frame
    _add_frame(pages, deck, question, scope or {}, research_mode, sources, claims, reportable_claims)
    _add_sub_questions(pages, deck, sub_questions or [], sources)

    if compact:
        _add_uncertainty(pages, deck, "3", t("uncertainty"), package, uncertainty)
    else:
        for index, section in enumerate(sections, 1):
            number = f"3.{index}"
            _add_prose(
                pages,
                deck,
                number,
                t("themes"),
                _clean(section.title),
                _section_text(section),
                BODY,
                f"theme:{index}",
                "sections[].synthesis",
            )
            has_findings = any(
                _clean(value) for value in (section.consensus, section.disagreements, section.implications)
            )
            if has_findings:
                _add_findings(pages, deck, number, section)
        _add_figures(pages, deck, research_figures or [], figure_observations or [])
        assessment = getattr(package, "cross_study_assessment", "") if package is not None else narrative
        _add_prose(
            pages,
            deck,
            "4",
            t("assessment_section"),
            t("assessment"),
            _paragraphs(assessment) or [t("no_assessment")],
            BODY,
            "assessment",
            "cross_study_assessment",
        )
        conclusion = getattr(package, "conclusion", "") if package is not None else ""
        _add_prose(
            pages,
            deck,
            "4",
            t("assessment_section"),
            t("conclusion"),
            _paragraphs(conclusion) or [t("no_conclusion")],
            BODY,
            "conclusion",
            "conclusion",
        )
        _add_uncertainty(pages, deck, "4", t("assessment_section"), package, uncertainty)

    # Appendices
    _add_method(pages, deck, coverage or {}, reportable_claims, run_id, package)
    _add_landscape(pages, deck, package, sections, figures_map)
    _add_source_catalog(pages, deck, sources, evidence_by_claim)
    _add_claim_register(pages, deck, reportable_claims, evidence_by_claim, source_numbers)
    pages.append(_Page(lambda slide, page: _draw_closing(deck, slide)))

    first_page = {}
    for number, page in enumerate(pages, 1):
        if page.toc_key and page.toc_key not in first_page:
            first_page[page.toc_key] = number
    for offset, chunk in enumerate(contents_chunks):
        pages[toc_index + offset] = _Page(
            lambda slide, page, chunk=chunk, offset=offset: _draw_contents(
                deck,
                slide,
                page,
                chunk,
                contents_style,
                first_page,
                first=offset == 0,
                last=offset == len(contents_chunks) - 1,
            )
        )

    for number, page in enumerate(pages, 1):
        page.draw(deck.slide(), number)

    return PresentationReportResult(document=deck.save(), figures=figures_map, citations=citations)


# -- cover, contents, closing ---------------------------------------------------------------


def _draw_cover(deck: _Deck, slide: Any, title: str, run_id: str) -> None:
    t = deck.t
    deck.rect(slide, "cover_panel", 0, 0, COVER_PANEL_W, SLIDE_H, RED)
    deck.picture(slide, "logo", deck.logo_white, 48, 54, 172.5, 116)
    date_style = TextStyle(18, bold=True, color=WHITE, leading=1.2)
    label_style = TextStyle(12, bold=True, color=WHITE, leading=1.2, tracking=1.0)
    date_y = SLIDE_H - 54 - date_style.line_height
    deck.text(slide, "cover_date_label", 48, date_y - 6 - label_style.line_height, 240, label_style.line_height, t("report_date"), label_style)
    deck.text(slide, "cover_date", 48, date_y, 240, date_style.line_height, datetime.now(UTC).strftime("%d.%m.%Y"), date_style)

    x, width = COVER_PANEL_W + 66, SLIDE_W - COVER_PANEL_W - 132
    title = _clean(title)
    meta = f"{t('run_id')} {run_id}"
    meta_h = text_height(meta, SMALL, width)
    chrome = EYEBROW.line_height + 18 + 18 + 2.25 + 18 + meta_h
    # The title is never cut: it steps down whole sizes until the cover block fits the page.
    style, title_h = _COVER_TITLE_STEPS[-1][0], 0.0
    for candidate, max_lines in _COVER_TITLE_STEPS:
        style, title_h = candidate, text_height(title, candidate, width)
        if line_count(title, candidate, width) <= max_lines and chrome + title_h <= SLIDE_H - 2 * TOP:
            break
    block = chrome + title_h
    y = max(TOP, (SLIDE_H - block) / 2)
    deck.text(slide, "cover_eyebrow", x, y, width, EYEBROW.line_height, t("report"), EYEBROW)
    y += EYEBROW.line_height + 18
    deck.text(slide, "cover_title", x, y, width, title_h, title, style)
    y += title_h + 18
    deck.rect(slide, "cover_rule", x, y, 48, 2.25, RED)
    y += 2.25 + 18
    deck.text(slide, "cover_run_id", x, y, width, meta_h, meta, SMALL)


def _draw_closing(deck: _Deck, slide: Any) -> None:
    deck.rect(slide, "cover_panel", 0, 0, COVER_PANEL_W, SLIDE_H, RED)
    deck.picture(slide, "logo", deck.logo_white, 48, 54, 172.5, 116)
    x, width = COVER_PANEL_W + 66, SLIDE_W - COVER_PANEL_W - 132
    thanks = TextStyle(48, bold=True, leading=1.1)
    foundation = TextStyle(16, color=GREY, leading=1.35)
    block = thanks.line_height + 18 + 2.25 + 18 + foundation.line_height
    y = (SLIDE_H - block) / 2
    deck.text(slide, "closing_title", x, y, width, thanks.line_height, deck.t("thanks"), thanks)
    y += thanks.line_height + 18
    deck.rect(slide, "closing_rule", x, y, 48, 2.25, RED)
    y += 2.25 + 18
    deck.text(slide, "closing_foundation", x, y, width, foundation.line_height, deck.t("foundation"), foundation)


_TOC_MAIN = TextStyle(16, leading=1.3)
_TOC_SUB_STYLES = (TextStyle(14, color=GREY, leading=1.3), TextStyle(12, color=GREY, leading=1.3))
_TOC_PAGE_W = 30.0
_TOC_APPENDIX_ROW = _TOC_MAIN.line_height + 12


def _contents_rows(deck: _Deck, compact: bool, sections: list[Any]) -> list[tuple[str, str, str, int]]:
    t = deck.t
    rows = [("1", t("summary"), "summary", 0), ("2", t("frame"), "frame", 0)]
    if compact:
        rows.append(("3", t("uncertainty"), "uncertainty", 0))
        return rows
    if sections:
        rows.append(("3", t("themes"), "theme:1", 0))
        rows.extend((f"3.{i}", _clean(s.title), f"theme:{i}", 1) for i, s in enumerate(sections, 1))
    rows.append(("4", t("toc_assessment"), "assessment", 0))
    return rows


def _contents_geometry(level: int, sub_style: TextStyle) -> tuple[TextStyle, float, float, float, float]:
    """Style, indent, marker width, text width and padding of a contents row."""
    if level == 0:
        return _TOC_MAIN, 0.0, 33.0, CONTENT_W - 33.0 - _TOC_PAGE_W - 12, 6.0
    return sub_style, 33.0, 30.0, CONTENT_W - 63.0 - _TOC_PAGE_W - 12, 2.5


def _contents_row_height(row: tuple[str, str, str, int], sub_style: TextStyle) -> float:
    style, _, _, text_w, pad = _contents_geometry(row[3], sub_style)
    return text_height(row[1], style, text_w) + 2 * pad


def _plan_contents(
    rows: list[tuple[str, str, str, int]],
) -> tuple[TextStyle, list[list[tuple[str, str, str, int]]]]:
    """Fit the contents on one slide, stepping theme entries down a size before using a second."""
    available = BODY_BOTTOM - TOP
    head = EYEBROW.line_height + 9
    appendix = 12 + EYEBROW.line_height + 9 + 2 * _TOC_APPENDIX_ROW
    for sub_style in _TOC_SUB_STYLES:
        if head + sum(_contents_row_height(row, sub_style) for row in rows) + appendix <= available:
            return sub_style, [rows]
    sub_style = _TOC_SUB_STYLES[-1]
    chunks: list[list[tuple[str, str, str, int]]] = [[]]
    used = head
    for row in rows:
        height = _contents_row_height(row, sub_style)
        if chunks[-1] and used + height > available:
            chunks.append([])
            used = head
        # An entry taller than a whole slide continues on the next one, split between lines;
        # the continuation carries no marker and no page number.
        while used + height > available:
            style, _, _, text_w, pad = _contents_geometry(row[3], sub_style)
            head_text, rest = fit_lines(row[1], style, text_w, available - used - 2 * pad)
            chunks[-1].append((row[0], head_text, row[2], row[3]))
            chunks.append([])
            used = head
            row = ("", rest, "", row[3])
            height = _contents_row_height(row, sub_style)
        chunks[-1].append(row)
        used += height
    if used + appendix > available:
        chunks.append([])
    return sub_style, chunks


def _draw_contents(
    deck: _Deck,
    slide: Any,
    page: int,
    rows: list[tuple[str, str, str, int]],
    sub_style: TextStyle,
    first_page: dict[str, int],
    first: bool,
    last: bool,
) -> None:
    t = deck.t
    deck.rect(slide, "rail", 0, 0, RAIL_W, SLIDE_H, RAIL)
    deck.rect(slide, "rail_bar", RAIL_X, TOP, 36, 3, RED)
    rail_title = t("toc") if first else f"{t('toc')}{t('continued')}"
    deck.text(
        slide,
        "rail_number",
        RAIL_X,
        TOP + 12,
        RAIL_INNER_W,
        text_height(rail_title, HEADING, RAIL_INNER_W),
        rail_title,
        HEADING,
    )
    deck.logo(slide)
    deck.footer(slide, page)

    y = TOP
    if rows:
        deck.text(slide, "toc_report", CONTENT_X, y, CONTENT_W, EYEBROW.line_height, t("toc_report"), EYEBROW)
        y += EYEBROW.line_height + 9
    for index, row in enumerate(rows, 1):
        marker, label, key, level = row
        style, indent, marker_w, text_w, pad = _contents_geometry(level, sub_style)
        height = text_height(label, style, text_w)
        if index == 1:
            deck.rect(slide, "toc_rule", CONTENT_X, y, CONTENT_W, 1.5, INK)
        elif level == 0:
            deck.rect(slide, f"toc_rule_{index}", CONTENT_X, y, CONTENT_W, 0.75, RULE)
        y += pad
        marker_style = replace(style, bold=level == 0, color=RED if level == 0 else MUTED)
        deck.text(slide, f"toc_marker_{index}", CONTENT_X + indent, y, marker_w, style.line_height, marker, marker_style)
        deck.text(slide, f"toc_entry_{index}", CONTENT_X + indent + marker_w, y, text_w, height, label, style)
        number = first_page.get(key)
        if number:
            deck.text(
                slide,
                f"toc_page_{index}",
                CONTENT_X + CONTENT_W - _TOC_PAGE_W,
                y,
                _TOC_PAGE_W,
                style.line_height,
                str(number),
                replace(style, color=MUTED),
                align=PP_ALIGN.RIGHT,
            )
        y += height + pad
    if not last:
        return

    y += 12 if rows else 0
    deck.text(slide, "toc_appendix", CONTENT_X, y, CONTENT_W, EYEBROW.line_height, t("toc_appendix"), EYEBROW)
    y += EYEBROW.line_height + 9
    appendices = [
        (t("appendix_a"), t("method_section"), "method"),
        (t("appendix_b"), t("landscape"), "landscape"),
        (t("appendix_c"), t("catalog"), "catalog"),
        (t("appendix_d"), t("claims"), "claims"),
    ]
    column_w = (CONTENT_W - 24) / 2
    marker_w = 42.0
    deck.rect(slide, "toc_appendix_rule", CONTENT_X, y, CONTENT_W, 1.5, INK)
    for index, (marker, label, key) in enumerate(appendices):
        column, row = index % 2, index // 2
        x = CONTENT_X + column * (column_w + 24)
        row_y = y + row * _TOC_APPENDIX_ROW
        if row:
            deck.rect(slide, f"toc_appendix_rule_{index}", x, row_y, column_w, 0.75, RULE)
        text_y = row_y + 6
        text_w = column_w - marker_w - _TOC_PAGE_W - 6
        deck.text(
            slide,
            f"toc_appendix_marker_{index}",
            x,
            text_y,
            marker_w,
            _TOC_MAIN.line_height,
            marker,
            replace(_TOC_MAIN, bold=True, color=RED),
        )
        deck.text(
            slide,
            f"toc_appendix_entry_{index}",
            x + marker_w,
            text_y,
            text_w,
            text_height(label, _TOC_MAIN, text_w),
            label,
            _TOC_MAIN,
        )
        number = first_page.get(key)
        if number:
            deck.text(
                slide,
                f"toc_appendix_page_{index}",
                x + column_w - _TOC_PAGE_W,
                text_y,
                _TOC_PAGE_W,
                _TOC_MAIN.line_height,
                str(number),
                replace(_TOC_MAIN, color=MUTED),
                align=PP_ALIGN.RIGHT,
            )


# -- report sections ------------------------------------------------------------------------


def _add_prose(
    pages: list[_Page],
    deck: _Deck,
    number: str,
    section_label: str,
    heading: str | None,
    paragraphs: list[str],
    style: TextStyle,
    toc_key: str,
    name: str,
) -> None:
    inline = bool(heading) and not _heading_fits(heading)
    top_heading = None if inline else heading
    _, height = _body_height(top_heading)
    flows = _heading_flows(heading) if inline else []
    flows += [Flow(text, style, CONTENT_W, gap_before=style.size * 0.75, name=name) for text in paragraphs]
    chunks = paginate(flows, height)
    for index, chunk in enumerate(chunks):
        title = top_heading
        if top_heading and index:
            title = f"{top_heading}{deck.t('continued')}"
        labels = _cited(*(piece.text for piece in chunk))

        def draw(slide: Any, page: int, chunk=chunk, title=title, labels=labels) -> None:
            deck.rail(slide, number, section_label)
            body_top = TOP
            if title:
                body_top = deck.heading(slide, title) + HEADING_GAP
            deck.draw_flows(slide, chunk, CONTENT_X, body_top)
            deck.footer(slide, page, deck.citation_note(labels))

        pages.append(_Page(draw, toc_key if index == 0 else ""))


def _add_frame(
    pages: list[_Page],
    deck: _Deck,
    question: str,
    scope: dict[str, Any],
    research_mode: str,
    sources: list[Any],
    claims: list[Any],
    reportable_claims: list[Any],
) -> None:
    t = deck.t
    turkish = deck.turkish
    heading = t("question_scope")
    _, height = _body_height(heading)
    question_style = TextStyle(20, leading=1.35)
    flows = [
        Flow(
            _clean(question) or "—",
            question_style,
            CONTENT_W,
            label=t("question"),
            label_style=EYEBROW,
            label_gap=7,
            pad_bottom=18,
            name="question",
        )
    ]
    rows: list[tuple[str, str]] = []
    start, end = _format_date(scope.get("start_date")), _format_date(scope.get("end_date"))
    if start != "—" or end != "—":
        rows.append((t("date_range"), f"{start} – {end}"))
    mode = _MODES.get(research_mode)
    rows.append((t("mode"), mode[0 if turkish else 1] if mode else _clean(research_mode) or "—"))
    intent = _clean(scope.get("query_intent"))
    if intent:
        rows.append((t("intent"), intent))
    families = Counter(_clean(getattr(source, "family", "")) for source in sources)
    families.pop("", None)
    source_text = f"{len(sources)} {'kaynak' if turkish else 'sources'}"
    if families:
        parts = [
            f"{count} {_FAMILIES.get(name, (name, name))[0 if turkish else 1]}"
            for name, count in families.most_common()
        ]
        source_text += f" · {', '.join(parts)}"
    rows.append((t("sources_row"), source_text))
    if claims:
        excluded = len(claims) - len(reportable_claims)
        claim_text = (
            f"{len(claims)} iddia · {len(reportable_claims)} raporlanabilir"
            + (f", {excluded} sentezden dışlandı" if excluded > 0 else "")
            if turkish
            else f"{len(claims)} claims · {len(reportable_claims)} reportable"
            + (f", {excluded} excluded from synthesis" if excluded > 0 else "")
        )
        rows.append((t("claims_row"), claim_text))
    label_w = 165.0
    for index, (label, value) in enumerate(rows):
        flows.append(
            Flow(
                value,
                BODY,
                CONTENT_W,
                pad_top=10,
                pad_bottom=10,
                marker=label,
                marker_width=label_w,
                strong_rule=index == 0,
                rule=index > 0,
                name=f"scope_{index + 1}",
            )
        )
    chunks = paginate(flows, height)
    for index, chunk in enumerate(chunks):
        title = heading if index == 0 else f"{heading}{t('continued')}"

        def draw(slide: Any, page: int, chunk=chunk, title=title) -> None:
            deck.rail(slide, "2", t("frame"))
            body_top = deck.heading(slide, title) + HEADING_GAP
            _draw_definition_rows(deck, slide, chunk, body_top)
            deck.footer(slide, page)

        pages.append(_Page(draw, "frame" if index == 0 else ""))


def _draw_definition_rows(deck: _Deck, slide: Any, pieces: list[Placed], top: float) -> None:
    """Flows whose marker is a grey row label rather than a red list number."""
    for index, piece in enumerate(pieces, 1):
        flow = piece.flow
        y = top + piece.top
        if flow.strong_rule or flow.rule:
            deck.rect(slide, f"{flow.name}_rule", CONTENT_X, y, flow.width, 1.5 if flow.strong_rule else 0.75, INK if flow.strong_rule else RULE)
        y += flow.pad_top
        if flow.label and flow.label_style is not None and not piece.continued:
            label_h = text_height(flow.label, flow.label_style, flow.width)
            deck.text(slide, f"{flow.name}_label", CONTENT_X, y, flow.width, label_h, flow.label, flow.label_style)
            y += label_h + flow.label_gap
        if flow.marker and not piece.continued:
            marker_style = TextStyle(14, color=GREY, leading=BODY.line_height / 14)
            deck.text(slide, f"{flow.name}_key", CONTENT_X, y, flow.marker_width - 12, BODY.line_height * 2, flow.marker, marker_style)
        height = text_height(piece.text, flow.style, flow.text_width)
        deck.text(slide, flow.name, CONTENT_X + flow.marker_width, y, flow.text_width, height, piece.text, flow.style)


def _add_sub_questions(pages: list[_Page], deck: _Deck, sub_questions: list[str], sources: list[Any]) -> None:
    t = deck.t
    heading = t("sub_questions")
    _, height = _body_height(heading)
    left_w = (CONTENT_W - COLUMN_GAP) * 7 / 12
    right_w = CONTENT_W - COLUMN_GAP - left_w
    questions = [_clean(item) for item in sub_questions if _clean(item)]
    if questions:
        flows = [
            Flow(
                text,
                LIST,
                left_w,
                pad_top=7,
                pad_bottom=7,
                marker=str(index),
                marker_width=21,
                strong_rule=index == 1,
                rule=index > 1,
                name=f"sub_question_{index}",
            )
            for index, text in enumerate(questions, 1)
        ]
    else:
        flows = [Flow(t("no_sub_questions"), LIST, left_w, name="sub_questions")]
    chunks = paginate(flows, height)

    near = [
        source
        for source in sources
        if (getattr(source, "metadata_json", {}) or {}).get("research_scope_role") in _SCOPE_ROLES
    ]
    near_top = EYEBROW.line_height + 8
    available = height - near_top - (SMALL.line_height + 12)
    near_items: list[tuple[str, str, float, float]] = []
    used = 0.0
    for source in near:
        title_text = _clean(getattr(source, "title", "")) or "—"
        role = _SCOPE_ROLES[source.metadata_json["research_scope_role"]][0 if deck.turkish else 1]
        title_h = text_height(title_text, SMALL_INK, right_w)
        item_h = 7 + title_h + 2 + SMALL.line_height + 7
        # A title is shown whole or not at all; the rest are counted below the list.
        if used + item_h > available:
            break
        near_items.append((title_text, role, used, title_h))
        used += item_h
    more = len(near) - len(near_items)

    for index, chunk in enumerate(chunks):
        title = heading if index == 0 else f"{heading}{t('continued')}"

        def draw(slide: Any, page: int, chunk=chunk, title=title, first=index == 0) -> None:
            deck.rail(slide, "2", t("frame"))
            body_top = deck.heading(slide, title) + HEADING_GAP
            deck.draw_flows(slide, chunk, CONTENT_X, body_top)
            if first:
                x = CONTENT_X + left_w + COLUMN_GAP
                deck.text(slide, "near_scope_label", x, body_top, right_w, EYEBROW.line_height, t("near"), EYEBROW)
                list_top = body_top + near_top
                if not near:
                    deck.text(
                        slide,
                        "near_scope_sources",
                        x,
                        list_top,
                        right_w,
                        text_height(t("no_near_scope"), SMALL_INK, right_w),
                        t("no_near_scope"),
                        SMALL_INK,
                    )
                for number, (title_text, role, offset, title_h) in enumerate(near_items, 1):
                    y = list_top + offset
                    deck.rect(slide, f"near_scope_{number}_rule", x, y, right_w, 1.5 if number == 1 else 0.75, INK if number == 1 else RULE)
                    deck.text(slide, f"near_scope_{number}", x, y + 7, right_w, title_h, title_text, SMALL_INK)
                    deck.text(slide, f"near_scope_{number}_role", x, y + 9 + title_h, right_w, SMALL.line_height, role, SMALL)
                if more:
                    y = list_top + used
                    note = (
                        f"+{more} çalışma daha · tam liste Word raporunda"
                        if deck.turkish
                        else f"+{more} more · full list in the Word report"
                    )
                    deck.rect(slide, "near_scope_more_rule", x, y, right_w, 0.75, RULE)
                    deck.text(slide, "near_scope_more", x, y + 6, right_w, SMALL.line_height, note, SMALL)
            deck.footer(slide, page)

        pages.append(_Page(draw, "frame_sub" if index == 0 else ""))


def _add_findings(pages: list[_Page], deck: _Deck, number: str, section: Any) -> None:
    t = deck.t
    heading = _clean(section.title)
    inline = not _heading_fits(heading)
    top_heading = None if inline else heading
    _, height = _body_height(top_heading)
    blocks = [
        ("consensus", t("consensus"), _clean(section.consensus) or t("no_consensus"), "sections[].consensus"),
        ("disagreements", t("disagreements"), _clean(section.disagreements) or t("no_disagreement"), "sections[].disagreements"),
        ("implications", t("implications"), _clean(section.implications) or t("no_implication"), "sections[].implications"),
    ]
    flows = [
        Flow(
            text,
            BODY,
            CONTENT_W,
            pad_top=10,
            pad_bottom=10,
            label=label,
            label_style=EYEBROW,
            label_gap=4,
            strong_rule=index == 0,
            rule=index > 0,
            name=name,
        )
        for index, (_, label, text, name) in enumerate(blocks)
    ]
    if inline:
        flows = _heading_flows(heading) + flows
    chunks = paginate(flows, height)
    for index, chunk in enumerate(chunks):
        title = None
        if top_heading:
            title = top_heading if index == 0 else f"{top_heading}{t('continued')}"
        labels = _cited(*(piece.text for piece in chunk))

        def draw(slide: Any, page: int, chunk=chunk, title=title, labels=labels) -> None:
            deck.rail(slide, number, t("themes"))
            body_top = deck.heading(slide, title) + HEADING_GAP if title else TOP
            deck.draw_flows(slide, chunk, CONTENT_X, body_top)
            deck.footer(slide, page, deck.citation_note(labels))

        pages.append(_Page(draw))


def _figures_to_render(research_figures: list[Any], observations: list[Any]) -> list[dict[str, Any]]:
    """The figures that get a slide: every exported research figure, then observed figures with an image.

    An observation without an exported image gets no slide, as in the Word report; run
    01M2FYR67BS5WXFMEVY2RT1EHP's fourth figure slide showed the template placeholder under a
    reading of an image the reader never saw.
    """
    by_hash = {obs.image_hash: obs for obs in observations if getattr(obs, "image_hash", None)}
    rendered: list[dict[str, Any]] = []
    for figure in research_figures:
        obs = by_hash.get(getattr(figure, "observation_hash", ""))
        labels = list(getattr(figure, "source_labels", []) or [])
        rendered.append(
            {
                "title": getattr(figure, "title", "") or (getattr(obs, "title", "") if obs else ""),
                "caption": getattr(figure, "caption", "") or (getattr(obs, "caption", "") if obs else ""),
                "attribution": getattr(figure, "attribution", "") or (getattr(obs, "source_title", "") if obs else ""),
                "rights": getattr(figure, "rights_statement", ""),
                "label": (labels[0] if labels else "") or (getattr(obs, "source_label", "") if obs else ""),
                "interpretation": getattr(obs, "selection_reason", "") if obs else "",
                "main_findings": list(getattr(obs, "main_findings", []) or []) if obs else [],
                "limitations": list(getattr(obs, "limitations", []) or []) if obs else [],
                "data": getattr(figure, "data", None),
            }
        )
    covered = {getattr(figure, "observation_hash", "") for figure in research_figures}
    for obs in observations:
        if getattr(obs, "image_hash", "") in covered:
            continue
        data = getattr(obs, "data", None) or getattr(obs, "image_bytes", None)
        if not data:
            continue
        rendered.append(
            {
                "title": getattr(obs, "title", "") or getattr(obs, "recommended_section", ""),
                "caption": getattr(obs, "caption", ""),
                "attribution": getattr(obs, "source_title", ""),
                "rights": "",
                "label": getattr(obs, "source_label", ""),
                "interpretation": getattr(obs, "selection_reason", ""),
                "main_findings": list(getattr(obs, "main_findings", []) or []),
                "limitations": list(getattr(obs, "limitations", []) or []),
                "data": data,
            }
        )
    return [item for item in rendered if item["data"]]


def _add_figures(pages: list[_Page], deck: _Deck, research_figures: list[Any], observations: list[Any]) -> None:
    t = deck.t
    title_top = TOP + EYEBROW.line_height + 7
    for number, figure in enumerate(_figures_to_render(research_figures, observations), 1):
        title = _FIGURE_NUMBER.sub("", _clean(figure["title"])) or (
            "Araştırma görseli" if deck.turkish else "Research figure"
        )
        label = _clean(figure["label"])
        eyebrow = f"{t('figure')} {number}" + (f" · {label}" if label else "")
        # An observation without an exported figure brings its raw "Figure 2: ..." caption.
        caption_parts = [_FIGURE_NUMBER.sub("", _clean(figure["caption"]))]
        attribution = re.sub(r"^(?:kaynak|source)\s*:\s*", "", _clean(figure["attribution"]), flags=re.IGNORECASE)
        if attribution and attribution != label:
            caption_parts.append(f"{t('source')}: {attribution}")
        rights = re.sub(r"^(?:telif|rights)\s*:\s*", "", _clean(figure["rights"]), flags=re.IGNORECASE)
        if rights:
            caption_parts.append(f"{t('rights')}: {rights}")
        caption = " · ".join(part for part in caption_parts if part)

        title_style = next(
            style
            for style, max_lines in _FIGURE_TITLE_STEPS
            if max_lines is None or line_count(title, style, CONTENT_W) <= max_lines
        )
        body_top = title_top + text_height(title, title_style, CONTENT_W) + 18
        # The image keeps its minimum height; a caption longer than the room left below it
        # continues on the next slide rather than covering the image.
        caption_room = BODY_BOTTOM - body_top - _FIGURE_IMAGE_MIN_H - 12
        shown, rest = caption, ""
        if caption and text_height(caption, SMALL, CONTENT_W) > caption_room:
            if caption_room < SMALL.line_height:
                shown, rest = "", caption
            else:
                shown, rest = fit_sentences(caption, SMALL, CONTENT_W, caption_room)
                if not shown:
                    shown, rest = fit_lines(caption, SMALL, CONTENT_W, caption_room)
        pages.append(
            _Page(
                lambda slide, page, figure=figure, title=title, style=title_style, eyebrow=eyebrow, shown=shown, label=label, body_top=body_top: _draw_figure(
                    deck, slide, page, figure, title, style, eyebrow, shown, label, body_top
                )
            )
        )
        if not rest:
            continue
        heading = f"{title}{t('continued')}"
        top = title_top + text_height(heading, title_style, CONTENT_W) + 18
        for chunk in paginate([Flow(rest, SMALL, CONTENT_W, name="figure.caption_attribution")], BODY_BOTTOM - top):
            pages.append(
                _Page(
                    lambda slide, page, chunk=chunk, heading=heading, style=title_style, eyebrow=eyebrow, label=label, top=top: _draw_figure_caption(
                        deck, slide, page, eyebrow, heading, style, chunk, label, top
                    )
                )
            )


def _draw_figure_heading(deck: _Deck, slide: Any, eyebrow: str, title: str, style: TextStyle) -> None:
    deck.rail(slide, "3", deck.t("figures_section"))
    deck.text(slide, "eyebrow", CONTENT_X, TOP, CONTENT_W, EYEBROW.line_height, eyebrow, EYEBROW)
    title_top = TOP + EYEBROW.line_height + 7
    deck.text(slide, "content_heading", CONTENT_X, title_top, CONTENT_W, text_height(title, style, CONTENT_W), title, style)


def _draw_figure_caption(
    deck: _Deck,
    slide: Any,
    page: int,
    eyebrow: str,
    heading: str,
    style: TextStyle,
    chunk: list[Placed],
    label: str,
    top: float,
) -> None:
    _draw_figure_heading(deck, slide, eyebrow, heading, style)
    deck.draw_flows(slide, chunk, CONTENT_X, top)
    deck.footer(slide, page, f"{deck.t('source')}  {label}" if label else "")


def _draw_figure(
    deck: _Deck,
    slide: Any,
    page: int,
    figure: dict[str, Any],
    title: str,
    title_style: TextStyle,
    eyebrow: str,
    caption: str,
    label: str,
    body_top: float,
) -> None:
    t = deck.t
    _draw_figure_heading(deck, slide, eyebrow, title, title_style)
    caption_h = text_height(caption, SMALL, CONTENT_W) if caption else 0.0
    box_h = max(_FIGURE_IMAGE_MIN_H, BODY_BOTTOM - body_top - (caption_h + 12 if caption else 0))
    image_w = (CONTENT_W - COLUMN_GAP) * 7 / 12
    side_x = CONTENT_X + image_w + COLUMN_GAP
    side_w = CONTENT_W - image_w - COLUMN_GAP

    deck.rect(slide, "figure_frame", CONTENT_X, body_top, image_w, box_h, RAIL)
    deck.picture(slide, "figure.asset", figure["data"], CONTENT_X + 12, body_top + 12, image_w - 24, box_h - 24)

    items = [_clean(item) for item in figure["main_findings"] if _clean(item)][:3]
    if not items and _clean(figure["interpretation"]):
        items = [_clean(figure["interpretation"])]
    limitation = next((_clean(item) for item in figure["limitations"] if _clean(item)), "")
    heading_h = EYEBROW.line_height + 8
    available = box_h - heading_h
    flows = [
        Flow(item, LIST, side_w, gap_before=8, marker="■", marker_width=14, name=f"figure.interpretation_{i}")
        for i, item in enumerate(items, 1)
    ]
    if limitation:
        flows.append(
            Flow(
                f"{t('limitation')}: {limitation}",
                TextStyle(12, color=GREY, leading=1.35),
                side_w,
                gap_before=10,
                pad_top=8,
                rule=True,
                name="figure.limitation",
            )
        )
    # Findings are the model's reading of the figure; keep the ones that fit whole, in order.
    while flows:
        fitted = paginate(flows, available)
        if len(fitted) == 1:
            break
        flows = flows[:-2] + flows[-1:] if limitation and len(flows) > 2 else flows[:-1]
    if items:
        deck.text(slide, "figure_reading_label", side_x, body_top, side_w, EYEBROW.line_height, t("implications"), EYEBROW)
    for piece in paginate(flows, available)[0] if flows else []:
        flow = piece.flow
        y = body_top + heading_h + piece.top
        if flow.rule:
            deck.rect(slide, f"{flow.name}_rule", side_x, y, side_w, 0.75, RULE)
        y += flow.pad_top
        if flow.marker:
            deck.text(slide, f"{flow.name}_marker", side_x, y + 1, flow.marker_width, flow.style.line_height, flow.marker, TextStyle(8, color=RED, leading=flow.style.line_height / 8))
        deck.text(slide, flow.name, side_x + flow.marker_width, y, flow.text_width, text_height(piece.text, flow.style, flow.text_width), piece.text, flow.style)
    if caption:
        deck.text(slide, "figure.caption_attribution", CONTENT_X, BODY_BOTTOM - caption_h, CONTENT_W, caption_h, caption, SMALL)
    deck.footer(slide, page, f"{t('source')}  {label}" if label else "")


def _add_uncertainty(
    pages: list[_Page], deck: _Deck, number: str, section_label: str, package: Any, fallback: str
) -> None:
    t = deck.t
    text = (getattr(package, "uncertainty", "") if package is not None else fallback) or ""
    heading = t("uncertainty")
    _, height = _body_height(heading)
    items = sentences(text)
    if items:
        flows = [
            Flow(
                item,
                BODY,
                CONTENT_W,
                pad_top=10,
                pad_bottom=10,
                marker=str(index),
                marker_width=24,
                strong_rule=index == 1,
                rule=index > 1,
                name=f"uncertainty_{index}",
            )
            for index, item in enumerate(items, 1)
        ]
    else:
        flows = [Flow(t("no_uncertainty"), BODY, CONTENT_W, name="uncertainty")]
    chunks = paginate(flows, height)
    for index, chunk in enumerate(chunks):
        title = heading if index == 0 else f"{heading}{t('continued')}"
        labels = _cited(*(piece.text for piece in chunk))

        def draw(slide: Any, page: int, chunk=chunk, title=title, labels=labels) -> None:
            deck.rail(slide, number, section_label)
            body_top = deck.heading(slide, title) + HEADING_GAP
            deck.draw_flows(slide, chunk, CONTENT_X, body_top)
            deck.footer(slide, page, deck.citation_note(labels))

        pages.append(_Page(draw, "uncertainty" if index == 0 else ""))


# -- appendices -----------------------------------------------------------------------------


def _add_method(
    pages: list[_Page],
    deck: _Deck,
    coverage: dict[str, Any],
    reportable_claims: list[Any],
    run_id: str,
    package: Any,
) -> None:
    t = deck.t
    turkish = deck.turkish
    unresolved = coverage.get("unresolved_major_claims", coverage.get("unresolved_primary_claims"))
    metrics = [
        ("Kaynak ailesi kapsamı" if turkish else "Source family coverage", _percent(coverage.get("source_family_coverage"), turkish)),
        ("Sorgu dalı kapsamı" if turkish else "Query branch coverage", _percent(coverage.get("query_branch_coverage"), turkish)),
        ("İddia denetim kapsamı" if turkish else "Claim audit coverage", _percent(coverage.get("claim_audit_coverage"), turkish)),
        ("Tahmini tamlık" if turkish else "Estimated completeness", _percent(coverage.get("estimated_completeness"), turkish)),
        ("Çözülmemiş ana iddia" if turkish else "Unresolved major claims", str(unresolved) if isinstance(unresolved, int) else "—"),
        ("Raporlanabilir iddia" if turkish else "Reportable claims", str(len(reportable_claims))),
    ]
    mode = getattr(package, "report_mode", "standard") if package is not None else "standard"
    meta = (
        f"{t('run_id')} {run_id} · {datetime.now(UTC).strftime('%d.%m.%Y %H:%M')} UTC · "
        f"pipeline v{REPORT_PIPELINE_VERSION} · {mode}"
    )

    def draw(slide: Any, page: int) -> None:
        deck.rail(slide, t("appendix_a"), t("method_section"))
        body_top = deck.heading(slide, t("method")) + HEADING_GAP
        column_w = (CONTENT_W - COLUMN_GAP) / 2
        steps = [
            Flow(step, LIST, column_w, pad_top=7, pad_bottom=7, marker=str(i), marker_width=18, strong_rule=i == 1, rule=i > 1, name=f"method_{i}")
            for i, step in enumerate(_METHOD_STEPS[turkish], 1)
        ]
        deck.draw_flows(slide, paginate(steps, BODY_BOTTOM - body_top)[0], CONTENT_X, body_top)
        x = CONTENT_X + column_w + COLUMN_GAP
        y = body_top
        header = SMALL_BOLD
        deck.text(slide, "metrics_header", x, y + 7, column_w - 60, header.line_height, t("metric"), header)
        deck.text(slide, "metrics_header_value", x + column_w - 60, y + 7, 60, header.line_height, t("value"), header, align=PP_ALIGN.RIGHT)
        y += 7 + header.line_height + 7
        deck.rect(slide, "metrics_rule", x, y, column_w, 1.5, INK)
        for index, (label, value) in enumerate(metrics, 1):
            deck.text(slide, f"metric_{index}", x, y + 8, column_w - 70, LIST.line_height, label, LIST)
            deck.text(slide, f"metric_{index}_value", x + column_w - 70, y + 8, 70, LIST.line_height, value, replace(LIST, bold=True), align=PP_ALIGN.RIGHT)
            y += 16 + LIST.line_height
            deck.rect(slide, f"metric_{index}_rule", x, y, column_w, 0.75, RULE)
        deck.footer(slide, page, meta)

    pages.append(_Page(draw, "method"))


def _add_landscape(
    pages: list[_Page], deck: _Deck, package: Any, sections: list[Any], figures_map: dict[str, bytes]
) -> None:
    t = deck.t
    profiles = list(getattr(package, "study_profiles", []) or []) if package is not None else []
    contributions = Counter(_clean(getattr(p, "contribution", "")) for p in profiles)
    contributions.pop("", None)
    cited = [cited_labels(section) for section in sections[:5]]
    matrix_rows = [
        (profile.source_label, [profile.source_label in labels for labels in cited])
        for profile in profiles
        if any(profile.source_label in labels for labels in cited)
    ]
    matrix_rows.sort(key=lambda row: (-sum(row[1]), _label_number(row[0])))
    landscape_image = figures_map.get("16a_research_contribution_landscape.png") or next(
        (v for k, v in figures_map.items() if "landscape" in k.lower() or "contribution" in k.lower()), None
    )
    map_image = figures_map.get("16b_theme_evidence_map.png") or next(
        (v for k, v in figures_map.items() if "theme" in k.lower() or "evidence_map" in k.lower()), None
    )

    def draw(slide: Any, page: int) -> None:
        deck.rail(slide, t("appendix_b"), t("landscape"))
        top = deck.heading(slide, t("landscape")) + HEADING_GAP
        height = BODY_BOTTOM - top
        matrix_w = 240.0
        chart_w = CONTENT_W - matrix_w - COLUMN_GAP
        caption_h = SMALL.line_height * 2
        if contributions:
            _draw_contribution_chart(deck, slide, contributions, CONTENT_X, top, chart_w, height - caption_h - 8)
        elif landscape_image:
            deck.picture(slide, "contribution_landscape", landscape_image, CONTENT_X, top, chart_w, height - caption_h - 8)
        else:
            deck.text(slide, "contribution_landscape", CONTENT_X, top, chart_w, LIST.line_height * 2, t("no_chart"), LIST)
        chart_caption = (
            "Araştırma katkısı türlerine göre kaynak sayısı" if deck.turkish else "Sources by research contribution type"
        )
        deck.text(slide, "landscape_caption", CONTENT_X, BODY_BOTTOM - caption_h, chart_w, caption_h, chart_caption, SMALL)
        x = CONTENT_X + chart_w + COLUMN_GAP
        if matrix_rows:
            _draw_theme_matrix(deck, slide, matrix_rows, len(cited), x, top, matrix_w, height)
        elif map_image:
            deck.picture(slide, "theme_evidence_map", map_image, x, top, matrix_w, height - caption_h - 8)
        else:
            deck.text(slide, "theme_evidence_map", x, top, matrix_w, LIST.line_height * 2, t("no_map"), LIST)
        deck.footer(slide, page)

    pages.append(_Page(draw, "landscape"))


def _draw_contribution_chart(
    deck: _Deck, slide: Any, contributions: Counter, x: float, y: float, width: float, height: float
) -> None:
    data = CategoryChartData()
    rows = contributions.most_common()
    data.categories = [label for label, _ in rows]
    data.add_series("Kaynak" if deck.turkish else "Sources", [count for _, count in rows])
    frame = slide.shapes.add_chart(XL_CHART_TYPE.BAR_CLUSTERED, Pt(x), Pt(y), Pt(width), Pt(height), data)
    frame.name = "contribution_landscape"
    chart = frame.chart
    chart.has_legend = False
    chart.has_title = False
    chart.font.name = FONT_NAME
    chart.font.size = Pt(12)
    chart.font.color.rgb = RGBColor.from_string(INK)
    plot = chart.plots[0]
    plot.gap_width = 55
    plot.has_data_labels = True
    labels = plot.data_labels
    labels.font.size = Pt(12)
    labels.font.name = FONT_NAME
    labels.number_format = "0"
    labels.number_format_is_linked = False
    labels.position = XL_LABEL_POSITION.OUTSIDE_END
    series = plot.series[0]
    series.format.fill.solid()
    series.format.fill.fore_color.rgb = RGBColor.from_string(RED)
    categories = chart.category_axis
    categories.reverse_order = True
    categories.has_major_gridlines = False
    categories.format.line.fill.background()
    categories.tick_labels.font.size = Pt(12)
    categories.tick_labels.font.name = FONT_NAME
    values = chart.value_axis
    values.visible = False
    values.has_major_gridlines = False


def _draw_theme_matrix(
    deck: _Deck,
    slide: Any,
    rows: list[tuple[str, list[bool]]],
    columns: int,
    x: float,
    y: float,
    width: float,
    height: float,
) -> None:
    label_w = 48.0
    column_w = (width - label_w) / max(1, columns)
    row_h = 16.0
    header_h = SMALL_BOLD.line_height + 6
    note_h = SMALL.line_height * 2 + 6
    capacity = max(1, int((height - header_h - note_h) // row_h))
    shown = rows[:capacity]
    deck.text(slide, "theme_matrix_header", x, y, label_w, SMALL_BOLD.line_height, "Kaynak" if deck.turkish else "Source", SMALL_BOLD)
    for column in range(columns):
        deck.text(
            slide,
            f"theme_matrix_column_{column + 1}",
            x + label_w + column * column_w,
            y,
            column_w,
            SMALL_BOLD.line_height,
            f"3.{column + 1}",
            SMALL_BOLD,
            align=PP_ALIGN.CENTER,
        )
    top = y + header_h
    deck.rect(slide, "theme_matrix_rule", x, top - 3, width, 1.5, INK)
    for index, (label, flags) in enumerate(shown):
        row_y = top + index * row_h
        deck.text(slide, f"theme_matrix_row_{index + 1}", x, row_y + 1, label_w, row_h, label, SMALL_INK)
        for column, active in enumerate(flags):
            deck.rect(
                slide,
                f"theme_matrix_cell_{index + 1}_{column + 1}",
                x + label_w + column * column_w + 3,
                row_y + 3,
                column_w - 6,
                row_h - 6,
                RED if active else RAIL,
            )
    if deck.turkish:
        note = "Sütunlar tema numaralarıdır; dolu hücre, kaynağın o temada atıf aldığını gösterir."
        if len(shown) < len(rows):
            note = f"En çok temaya katkı veren {len(shown)} / {len(rows)} kaynak. " + note
    else:
        note = "Columns are theme numbers; a filled cell means the theme cites the source."
        if len(shown) < len(rows):
            note = f"{len(shown)} of {len(rows)} sources, most themes first. " + note
    note_y = top + len(shown) * row_h + 6
    deck.text(slide, "theme_matrix_note", x, note_y, width, text_height(note, SMALL, width), note, SMALL)


def _clean_title(title: Any) -> str:
    cleaned = _clean(title)
    stripped = _TITLE_SUFFIX.sub("", cleaned).strip()
    return stripped or cleaned


def _domain(url: Any) -> str:
    host = urlparse(str(url or "")).netloc.lower()
    return host.removeprefix("www.")


def _table_page(
    deck: _Deck,
    slide: Any,
    columns: list[tuple[str, float, Any]],
    rows: list[list[str]],
    top: float,
    bottom: float,
    name: str,
) -> int:
    """Draw the header and as many whole rows as fit; return how many rows were drawn."""
    flexible = CONTENT_W - sum(width for _, width, _ in columns if width) - 12 * (len(columns) - 1)
    widths = [width or flexible for _, width, _ in columns]
    x_positions = []
    x = CONTENT_X
    for width in widths:
        x_positions.append(x)
        x += width + 12
    y = top
    for index, ((label, _, align), width, x) in enumerate(zip(columns, widths, x_positions, strict=True)):
        deck.text(slide, f"{name}_header_{index + 1}", x, y + 6, width, SMALL_BOLD.line_height, label, SMALL_BOLD, align=align)
    y += 6 + SMALL_BOLD.line_height + 6
    deck.rect(slide, f"{name}_header_rule", CONTENT_X, y, CONTENT_W, 1.5, INK)
    drawn = 0
    for row_index, row in enumerate(rows, 1):
        styles = [SMALL_BOLD if column == 0 else SMALL_INK for column in range(len(columns))]
        height = max(text_height(value, style, width) for value, style, width in zip(row, styles, widths, strict=True))
        if y + 6 + height + 6 > bottom:
            break
        for column, (value, style, width, x) in enumerate(zip(row, styles, widths, x_positions, strict=True)):
            deck.text(slide, f"{name}_{row_index}_{column + 1}", x, y + 6, width, height, value, style, align=columns[column][2])
        y += 6 + height + 6
        deck.rect(slide, f"{name}_{row_index}_rule", CONTENT_X, y, CONTENT_W, 0.75, RULE)
        drawn += 1
    return drawn


def _add_source_catalog(
    pages: list[_Page], deck: _Deck, sources: list[Any], evidence_by_claim: dict[str, list[tuple[Any, Any]]]
) -> None:
    t = deck.t
    turkish = deck.turkish
    counts = _source_evidence_counts(sources, evidence_by_claim)
    numbered = list(enumerate(sources, 1))
    numbered.sort(key=lambda item: (-counts.get(str(item[1].id), 0), item[0]))
    rows = []
    for number, source in numbered:
        family = _clean(getattr(source, "family", ""))
        kind = _FAMILY_SHORT.get(family, (family, family))[0 if turkish else 1] if family else "—"
        rows.append(
            [
                f"S{number:02d}",
                _clean_title(getattr(source, "title", "")) or "—",
                _domain(getattr(source, "url", "")) or "—",
                kind,
                str(counts.get(str(source.id), 0)),
            ]
        )
    columns = [
        ("No", 36.0, None),
        ("Başlık" if turkish else "Title", 0.0, None),
        ("Alan adı" if turkish else "Domain", 144.0, None),
        ("Tür" if turkish else "Type", 57.0, None),
        ("İddia" if turkish else "Claims", 36.0, PP_ALIGN.RIGHT),
    ]
    note = (
        "Tam liste: Word raporu Ek C ve 05_source_catalog.csv"
        if turkish
        else "Full list: Word report Appendix C and 05_source_catalog.csv"
    )

    def draw(slide: Any, page: int) -> None:
        deck.rail(slide, t("appendix_c"), t("catalog"))
        heading_bottom = deck.heading(slide, t("catalog"), width=CONTENT_W - 220)
        top = heading_bottom + HEADING_GAP - 6
        if not rows:
            deck.text(slide, "catalog_empty", CONTENT_X, top, CONTENT_W, LIST.line_height, t("no_sources"), LIST)
            deck.footer(slide, page)
            return
        drawn = _table_page(deck, slide, columns, rows, top, BODY_BOTTOM, "catalog")
        tag = (
            f"İddia sayısına göre ilk {drawn} / {len(rows)} kaynak"
            if turkish
            else f"Top {drawn} of {len(rows)} sources by claims"
        )
        deck.text(slide, "catalog_tag", CONTENT_X + CONTENT_W - 220, TOP + 8, 220, SMALL.line_height, tag, SMALL, align=PP_ALIGN.RIGHT)
        deck.footer(slide, page, note if drawn < len(rows) else "")

    pages.append(_Page(draw, "catalog"))


def _add_claim_register(
    pages: list[_Page],
    deck: _Deck,
    reportable_claims: list[Any],
    evidence_by_claim: dict[str, list[tuple[Any, Any]]],
    source_numbers: dict[Any, int],
) -> None:
    t = deck.t
    turkish = deck.turkish
    rows = []
    for number, claim in enumerate(reportable_claims, 1):
        status = _clean(getattr(claim, "status", ""))
        audit = getattr(claim, "audit", {}) or {}
        sources = _claim_sources(str(getattr(claim, "id", "")), evidence_by_claim, source_numbers)
        rows.append(
            [
                f"C{number:02d}",
                _clean(getattr(claim, "text", "") or getattr(claim, "claim", "")) or "—",
                _STATUSES.get(status, (status, status))[0 if turkish else 1] or "—",
                _decimal(getattr(claim, "confidence", None), turkish),
                _decimal(audit.get("question_relevance"), turkish),
                ", ".join(f"S{n:02d}" for n in sources) or "—",
            ]
        )
    columns = [
        ("No", 36.0, None),
        ("İddia" if turkish else "Claim", 0.0, None),
        ("Durum" if turkish else "Status", 84.0, None),
        ("Güven" if turkish else "Conf.", 39.0, PP_ALIGN.RIGHT),
        ("İlgi" if turkish else "Rel.", 33.0, PP_ALIGN.RIGHT),
        ("Kaynak" if turkish else "Sources", 57.0, PP_ALIGN.RIGHT),
    ]
    note = (
        "Tam kayıt: Word raporu Ek D ve 04_claim_ledger.jsonl"
        if turkish
        else "Full register: Word report Appendix D and 04_claim_ledger.jsonl"
    )

    def draw(slide: Any, page: int) -> None:
        deck.rail(slide, t("appendix_d"), t("claims"))
        heading_bottom = deck.heading(slide, t("claims"), width=CONTENT_W - 220)
        top = heading_bottom + HEADING_GAP - 6
        if not rows:
            deck.text(slide, "claims_empty", CONTENT_X, top, CONTENT_W, LIST.line_height, t("no_claims"), LIST)
            deck.footer(slide, page)
            return
        drawn = _table_page(deck, slide, columns, rows, top, BODY_BOTTOM, "claims")
        tag = f"İlk {drawn} / {len(rows)} iddia" if turkish else f"First {drawn} of {len(rows)} claims"
        deck.text(slide, "claims_tag", CONTENT_X + CONTENT_W - 220, TOP + 8, 220, SMALL.line_height, tag, SMALL, align=PP_ALIGN.RIGHT)
        deck.footer(slide, page, note if drawn < len(rows) else "")

    pages.append(_Page(draw, "claims"))
