#!/usr/bin/env python3
"""List every slide's boxes -- name, position, size, text -- and flag the ones that collide.

The presentation polisher copies this file into the agent's working folder as
``outline.py``::

    python outline.py sunum.pptx                   # every slide, as JSON
    python outline.py cilali.pptx --check          # only problems: overlaps, off-slide boxes
    python outline.py cilali.pptx --slides 3,7-9 --text 400

Not ``inspect.py``: a file of that name in the working folder would shadow the standard
library module for every script run there. The polishing agent used to spend its first
minutes writing inspectors of its own, several of which died on the Windows console
encoding (measured 2026-09-16); this one always writes UTF-8.
"""

from __future__ import annotations

import argparse
import json
import sys
from pathlib import Path
from typing import Any

import pptx
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.util import Emu

#: Boxes may touch; they overlap only past this many points in both directions. The
#: generated template's scope keys sit 4 pt into each other by design.
_TOLERANCE_PT = 5.0
#: Template chrome in the side rail and footer, which sits apart from the content by design.
_CHROME_PREFIXES = ("rail", "logo", "footer", "slide_number", "brand")


def _pt(value: int | None) -> float:
    return round(Emu(value or 0).pt, 1)


def _slide_numbers(spec: str | None, count: int) -> list[int]:
    if not spec:
        return list(range(1, count + 1))
    numbers: set[int] = set()
    for part in spec.split(","):
        part = part.strip()
        if not part:
            continue
        if "-" in part:
            start, end = part.split("-", 1)
            numbers.update(range(int(start), int(end) + 1))
        else:
            numbers.add(int(part))
    return sorted(number for number in numbers if 1 <= number <= count)


def _describe(shape: Any, text_limit: int) -> dict[str, Any]:
    try:
        kind = shape.shape_type
        kind_name = getattr(kind, "name", str(kind))
    except NotImplementedError:
        kind, kind_name = None, "UNKNOWN"
    item: dict[str, Any] = {
        "name": shape.name,
        "kind": kind_name,
        "left": _pt(shape.left),
        "top": _pt(shape.top),
        "width": _pt(shape.width),
        "height": _pt(shape.height),
    }
    if kind == MSO_SHAPE_TYPE.PICTURE:
        item["image"] = True
    if shape.has_text_frame and shape.text_frame.text.strip():
        text = shape.text_frame.text
        item["text"] = text if len(text) <= text_limit else text[:text_limit] + "…"
        item["chars"] = len(text)
        sizes = sorted(
            {
                run.font.size.pt
                for paragraph in shape.text_frame.paragraphs
                for run in paragraph.runs
                if run.font.size is not None
            }
        )
        if sizes:
            item["font_pt"] = sizes
    if getattr(shape, "has_table", False) and shape.has_table:
        table = shape.table
        item["table"] = {
            "rows": len(table.rows),
            "columns": len(table.columns),
            "first_row": [cell.text[:40] for cell in table.rows[0].cells],
        }
    return item


def _issues(boxes: list[dict[str, Any]], width: float, height: float) -> list[str]:
    issues = []
    for box in boxes:
        if (
            box["left"] < -_TOLERANCE_PT
            or box["top"] < -_TOLERANCE_PT
            or box["left"] + box["width"] > width + _TOLERANCE_PT
            or box["top"] + box["height"] > height + _TOLERANCE_PT
        ):
            issues.append(f"off-slide: {box['name']}")
    content = [
        box
        for box in boxes
        if (box.get("text") or box.get("image") or box.get("table"))
        and not box["name"].lower().startswith(_CHROME_PREFIXES)
    ]
    for index, first in enumerate(content):
        for second in content[index + 1 :]:
            across = min(first["left"] + first["width"], second["left"] + second["width"]) - max(
                first["left"], second["left"]
            )
            down = min(first["top"] + first["height"], second["top"] + second["height"]) - max(
                first["top"], second["top"]
            )
            if across > _TOLERANCE_PT and down > _TOLERANCE_PT:
                issues.append(
                    f"overlap: {first['name']} <-> {second['name']} ({across:.0f} x {down:.0f} pt)"
                )
    return issues


def outline(deck: Path, *, slides: str | None, text_limit: int) -> dict[str, Any]:
    presentation = pptx.Presentation(str(deck))
    width, height = _pt(presentation.slide_width), _pt(presentation.slide_height)
    result: dict[str, Any] = {
        "deck": deck.name,
        "slide_count": len(presentation.slides),
        "slide_width_pt": width,
        "slide_height_pt": height,
        "slides": [],
    }
    all_slides = list(presentation.slides)
    for number in _slide_numbers(slides, len(all_slides)):
        slide = all_slides[number - 1]
        boxes = [_describe(shape, text_limit) for shape in slide.shapes]
        result["slides"].append(
            {
                "number": number,
                "layout": slide.slide_layout.name,
                "boxes": boxes,
                "issues": _issues(boxes, width, height),
            }
        )
    return result


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description="Outline a deck's slides and flag colliding boxes.")
    parser.add_argument("deck", type=Path, help="presentation, e.g. sunum.pptx or cilali.pptx")
    parser.add_argument("--slides", help="slide numbers, e.g. 3,7-9 (default: all)")
    parser.add_argument("--text", type=int, default=300, help="characters of text per box")
    parser.add_argument("--check", action="store_true", help="print only slides with problems")
    args = parser.parse_args(argv)
    # The Windows console code page cannot print Turkish text; JSON goes out as UTF-8.
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(encoding="utf-8")
    if not args.deck.is_file():
        print(f"not found: {args.deck}", file=sys.stderr)
        return 2
    try:
        result = outline(args.deck, slides=args.slides, text_limit=args.text)
    except Exception as exc:  # noqa: BLE001 - python-pptx raises several parser exceptions
        print(f"cannot read {args.deck.name}: {exc}", file=sys.stderr)
        return 1
    if args.check:
        problems = [
            {"number": slide["number"], "issues": slide["issues"]}
            for slide in result["slides"]
            if slide["issues"]
        ]
        result = {
            "deck": result["deck"],
            "slide_count": result["slide_count"],
            "slides_with_problems": len(problems),
            "problems": problems,
        }
    print(json.dumps(result, ensure_ascii=False, indent=1))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
