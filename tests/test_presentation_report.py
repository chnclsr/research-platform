import io
import re
import zipfile
from types import SimpleNamespace

import pptx
from pptx.oxml.ns import qn

from research_platform.presentation_layout import TextStyle, text_height
from research_platform.presentation_report import (
    COVER_PANEL_W,
    FOOTER_TOP,
    LOGO_COLOR,
    LOGO_WHITE,
    PRESENTATION_REPORT_FALLBACK,
    TEMPLATE_PATH,
    build_presentation_report,
    presentation_report_name,
)
from research_platform.report_synthesis import (
    _REPORT_WITHOUT_SUMMARY,
    StudyProfile,
    SynthesisPackage,
    SynthesisSection,
)


def _minimal_report_inputs():
    source1 = SimpleNamespace(
        id="src-1",
        title="Sağlıkta Yapay Zeka Uygulamaları",
        url="https://example.org/s1",
        connector_id="pubmed",
        metadata_json={"year": 2024, "type": "Makale", "research_scope_role": "primary_in_scope"},
    )
    claim1 = SimpleNamespace(
        id="claim-1",
        text="Yapay zeka modelleri tanı doğruluğunu artırır.",
        status="supported",
        confidence=0.95,
        audit={"question_relevance": 0.92},
    )
    link1 = SimpleNamespace(
        direction="supports",
        quote="Yapay zeka tanı sürecini hızlandırır.",
        location={"section_path": "Abstract"},
        entailment_score=0.95,
    )
    return {
        "run_id": "test-run-12345678",
        "title": "Sağlıkta Yapay Zeka",
        "question": "Yapay zeka tanı süreçlerini nasıl etkiler?",
        "language": "tr",
        "coverage": {
            "source_family_coverage": 0.90,
            "query_branch_coverage": 0.85,
            "claim_audit_coverage": 0.95,
            "estimated_completeness": 0.0159,
            "unresolved_major_claims": 0,
        },
        "sources": [source1],
        "claims": [claim1],
        "reportable_claims": [claim1],
        "evidence_by_claim": {"claim-1": [(link1, source1)]},
        "executive_summary": "Yapay zeka tanı süreçlerinde etkilidir [S01].",
        "narrative": "Kanıt sentezi detayları.",
        "uncertainty": "Veri kalitesi ve heterojenlik belirsizlik yaratmaktadır.",
        "scope": {
            "start_date": "2020",
            "end_date": "2024",
            "query_intent": "Klinik yapay zeka literatürü",
        },
        "sub_questions": ["Tanı performansı nedir?", "İş akışına etkisi nedir?"],
    }


def _package(sections, **overrides):
    values = {
        "executive_summary": "Özet cümlesi [S01].",
        "sections": sections,
        "cross_study_assessment": "Değerlendirme [S01].",
        "conclusion": "Sonuç [S01].",
        "uncertainty": "Belirsizlik [S01].",
        "study_profiles": [],
        "generated_by_llm": True,
        "report_mode": "standard",
    }
    values.update(overrides)
    return SynthesisPackage(**values)


def _deck(inputs):
    return pptx.Presentation(io.BytesIO(build_presentation_report(**inputs).document))


def _shape(slide, name):
    return next((shape for shape in slide.shapes if shape.name == name), None)


def _text(slide, name):
    shape = _shape(slide, name)
    return shape.text_frame.text if shape is not None and shape.has_text_frame else ""


def _texts(slide):
    return [shape.text_frame.text for shape in slide.shapes if shape.has_text_frame]


def _deck_text(prs):
    return "\n".join(text for slide in prs.slides for text in _texts(slide))


def _headings(prs):
    return [_text(slide, "content_heading") for slide in prs.slides]


def _contents(prs):
    """(marker, entry, page) rows of the contents slides."""
    rows = []
    for slide in prs.slides:
        if not _text(slide, "rail_number").startswith("İçindekiler"):
            continue
        for shape in slide.shapes:
            match = re.fullmatch(r"toc_(appendix_)?entry_(\d+)", shape.name)
            if match:
                prefix = f"toc_{match.group(1) or ''}"
                index = match.group(2)
                rows.append(
                    (
                        _text(slide, f"{prefix}marker_{index}"),
                        shape.text_frame.text,
                        _text(slide, f"{prefix}page_{index}"),
                    )
                )
    return rows


def _png(color, size=(640, 480)):
    from PIL import Image

    buffer = io.BytesIO()
    Image.new("RGB", size, color=color).save(buffer, format="PNG")
    return buffer.getvalue()


def _broken_image_links(prs):
    broken = []
    for number, slide in enumerate(prs.slides, 1):
        rels = slide.part.rels
        for element in slide._element.iter():
            if not isinstance(element.tag, str):
                continue
            rid = element.get(qn("r:embed"))
            if rid and rid not in rels:
                broken.append((number, rid))
    return broken


LONG_SENTENCES = [
    (
        "Akciğer bilgisayarlı tomografi taramalarında yapay zeka modellerinin tanısal doğruluğu, derin "
        "evrişimli sinir ağları ve transformatör tabanlı mimarilerle belirgin biçimde artmıştır [S01]."
    ),
    (
        "Çok merkezli kohortlarda 3 mm altındaki nodüllerin yakalanma hassasiyeti %95 üzerine çıkmış ve "
        "radyologların gözden kaçırma oranı önemli ölçüde düşürülmüştür [S02]."
    ),
    (
        "Malignite derecelendirmesinde radyomik özniteliklerin entegrasyonu ROC-AUC skorunu 0.93 seviyesine "
        "yükseltmiştir [S03]."
    ),
    (
        "Bununla birlikte DICOM standartlarındaki farklılıklar ve rekonstrüksiyon filtreleri modellerin "
        "genellenebilirliğinde sapmalara yol açabilmektedir [S04]."
    ),
]
LONG_TITLE = (
    "Yapay zeka tarafından oluşturulan BT radyoloji raporları klinik doğruluk, olgusal doğruluk ve dil "
    "kalitesi açısından nasıl değerlendirilir ve bu değerlendirmeler hangi kanıta dayanır?"
)


def _long_inputs():
    """A run the size of 01M2FYR67BS5WXFMEVY2RT1EHP: long themes, many sources and claims."""
    inputs = _minimal_report_inputs()
    sources = [
        SimpleNamespace(
            id=f"src-{i}",
            title=f"Automatic radiology report generation with deep learning, part {i}: a comprehensive "
            "review of methods and advances | Artificial Intelligence Review | Springer Nature Link",
            url=f"https://www.onlinelibrary.wiley.com/doi/full/10.1002/example.{i}",
            connector_id="agentsearch_web",
            family="web" if i % 5 else "code_data",
            metadata_json={"research_scope_role": "near_scope" if i % 4 == 0 else "primary_in_scope"},
        )
        for i in range(1, 41)
    ]
    claims = [
        SimpleNamespace(
            id=f"claim-{i}",
            text="The most recent generation of systems aims to combine volumetric medical imaging with "
            "language and generate clinically meaningful textual outputs, including reports. " * (1 + i % 3),
            status="qualified" if i % 2 else "supported",
            confidence=0.9,
            audit={"question_relevance": 0.8},
        )
        for i in range(1, 31)
    ]
    evidence = {
        f"claim-{i}": [(SimpleNamespace(direction="supports"), sources[(i + j) % 40]) for j in range(3)]
        for i in range(1, 31)
    }
    synthesis = " ".join(LONG_SENTENCES * 6)
    sections = [
        SynthesisSection(
            title=f"{LONG_TITLE} ({index})",
            synthesis=f"{synthesis}\n\n{synthesis}",
            consensus=" ".join(LONG_SENTENCES * 2),
            disagreements=LONG_SENTENCES[3],
            implications=" ".join(LONG_SENTENCES * 2),
            source_ids=["src-1"],
            claim_ids=["claim-1"],
        )
        for index in range(1, 6)
    ]
    profiles = [
        StudyProfile(
            source_id=f"src-{i}",
            source_label=f"S{i:02d}",
            title=f"Kaynak {i}",
            contribution=["Yöntem geliştirme", "Veri seti / benchmark", "Diğer katkı"][i % 3],
            evidence_design="review",
        )
        for i in range(1, 41)
    ]
    inputs.update(
        {
            "title": f"{LONG_TITLE} {LONG_TITLE}",
            "question": LONG_TITLE * 3,
            "sources": sources,
            "claims": claims,
            "reportable_claims": claims,
            "evidence_by_claim": evidence,
            "sub_questions": [LONG_TITLE] * 7,
            "synthesis_package": _package(
                sections,
                executive_summary=" ".join(LONG_SENTENCES * 5),
                cross_study_assessment=" ".join(LONG_SENTENCES * 8),
                conclusion=" ".join(LONG_SENTENCES * 3),
                uncertainty=" ".join(LONG_SENTENCES * 3),
                study_profiles=profiles,
            ),
            "research_figures": [
                SimpleNamespace(
                    title="Şekil 2: Hacimsel BT görüntülerinin transvers kesitleri ve karşılık gelen raporlar",
                    caption="Figure 2: Dataset examples. " + LONG_SENTENCES[0] * 2,
                    attribution="Kaynak: Automatically Generating Narrative-Style Radiology Reports",
                    rights_statement="Telif: Kaynak metadata kaydında açık lisans bilgisi bulunamadı.",
                    source_labels=["S47"],
                    data=_png("navy", (1600, 500)),
                    observation_hash="wide",
                ),
                SimpleNamespace(
                    title="Dikey figür",
                    caption="Şekil 3",
                    attribution="S48",
                    rights_statement="",
                    source_labels=["S48"],
                    data=_png("teal", (500, 1400)),
                    observation_hash="tall",
                ),
            ],
            "figure_observations": [
                SimpleNamespace(
                    image_hash="wide",
                    selection_reason="Veri kümesi örneği",
                    main_findings=[LONG_SENTENCES[0], LONG_SENTENCES[1], LONG_SENTENCES[2]],
                    limitations=[LONG_SENTENCES[3]],
                )
            ],
        }
    )
    return inputs


def test_presentation_report_name():
    assert presentation_report_name("ai_in_lung_ct") == "16_ai_in_lung_ct_report.pptx"
    assert presentation_report_name(None) == PRESENTATION_REPORT_FALLBACK
    assert presentation_report_name("") == PRESENTATION_REPORT_FALLBACK
    assert presentation_report_name("sağlık_yapay_zekâ").endswith(".pptx")


def test_cover_uses_the_report_title_and_the_frame_slide_the_question():
    inputs = _minimal_report_inputs()
    inputs["title"] = "Kısa Okuyucu Başlığı"
    inputs["question"] = "Uzun ve karmaşık orijinal araştırma sorusu?"
    prs = _deck(inputs)

    assert _text(prs.slides[0], "cover_title") == "Kısa Okuyucu Başlığı"
    frame = next(slide for slide in prs.slides if _text(slide, "content_heading") == "Ana soru ve kapsam")
    assert "Uzun ve karmaşık orijinal araştırma sorusu?" in _texts(frame)


def _all_shapes_on_the_page(prs):
    for number, slide in enumerate(prs.slides, 1):
        for shape in slide.shapes:
            assert shape.top.pt >= -0.5, (number, shape.name)
            assert shape.top.pt + shape.height.pt <= 540.5, (number, shape.name)
            assert shape.left.pt + shape.width.pt <= 960.5, (number, shape.name)


def test_a_long_cover_title_stays_whole_and_clear_of_the_logo():
    """Run 01M2FYR67BS5WXFMEVY2RT1EHP's cover title grew over the logo baked into the image."""
    # A report title is at most 120 characters and a protocol title 300; 1,000 leaves margin.
    for length in (len(LONG_TITLE) * 2, 1000):
        inputs = _minimal_report_inputs()
        inputs["title"] = " ".join(LONG_SENTENCES * 4)[:length].strip()
        prs = _deck(inputs)
        cover = prs.slides[0]

        title = _shape(cover, "cover_title")
        logo = _shape(cover, "logo")
        assert title.text_frame.text == inputs["title"]
        assert title.left.pt >= COVER_PANEL_W
        assert logo.left.pt + logo.width.pt <= COVER_PANEL_W
        _all_shapes_on_the_page(prs)


def test_an_oversized_theme_title_flows_with_its_text_instead_of_failing():
    title = " ".join(LONG_SENTENCES * 2)
    inputs = _minimal_report_inputs()
    inputs["synthesis_package"] = _package(
        [SynthesisSection(title=title, synthesis=" ".join(LONG_SENTENCES), consensus="Ortak yön [S01].")]
    )
    prs = _deck(inputs)

    _all_shapes_on_the_page(prs)
    theme = [slide for slide in prs.slides if _text(slide, "rail_number") == "3.1"]
    # Too tall to repeat above every continuation: the synthesis and the findings each open
    # with the whole title, and the slides that continue them carry the theme number only.
    opening = [slide for slide in theme if _shape(slide, "content_heading") is not None]
    assert len(opening) == 2
    for slide in opening:
        heading = " ".join(shape.text_frame.text for shape in slide.shapes if shape.name.startswith("content_heading"))
        assert heading == title
    synthesis = " ".join(
        shape.text_frame.text for slide in theme for shape in slide.shapes if shape.name.startswith("sections[].synthesis")
    )
    assert synthesis == " ".join(LONG_SENTENCES)


def test_a_theme_title_longer_than_a_slide_still_stays_on_the_page():
    # Sub-question titles carry no length limit in the protocol schema.
    title = " ".join(LONG_SENTENCES * 12)
    inputs = _minimal_report_inputs()
    inputs["synthesis_package"] = _package(
        [SynthesisSection(title=title, synthesis=" ".join(LONG_SENTENCES), consensus="Ortak yön [S01].")]
    )
    prs = _deck(inputs)

    _all_shapes_on_the_page(prs)
    entries = " ".join(entry for marker, entry, _ in _contents(prs) if marker in ("3.1", ""))
    assert entries == title


def test_a_long_figure_caption_continues_instead_of_covering_the_image():
    caption = " ".join(LONG_SENTENCES * 6)
    inputs = _minimal_report_inputs()
    inputs["research_figures"] = [
        SimpleNamespace(
            title="Uzun açıklamalı figür",
            caption=caption,
            attribution="",
            rights_statement="",
            data=_png("red"),
            observation_hash="h",
        )
    ]
    prs = _deck(inputs)

    _all_shapes_on_the_page(prs)
    slides = [slide for slide in prs.slides if _text(slide, "eyebrow") == "FİGÜR 1"]
    assert len(slides) > 1
    frame = _shape(slides[0], "figure_frame")
    assert frame.height.pt >= 150
    shown_caption = _shape(slides[0], "figure.caption_attribution")
    if shown_caption is not None:
        assert frame.top.pt + frame.height.pt <= shown_caption.top.pt
    assert all(_text(slide, "content_heading").endswith("(devam)") for slide in slides[1:])
    shown = " ".join(
        shape.text_frame.text
        for slide in slides
        for shape in slide.shapes
        if shape.name.startswith("figure.caption_attribution")
    )
    assert shown == caption


def test_compact_report_suppresses_thematic_slides():
    inputs = _minimal_report_inputs()
    inputs["synthesis_package"] = _package(
        [SynthesisSection(title="Kanıt özeti", synthesis="Seyrek kanıt [S01].", source_ids=["src-1"])],
        executive_summary="Seyrek kanıt yalnız sınırlı bir sonuç destekler [S01].",
        cross_study_assessment="",
        conclusion="",
        uncertainty="Kanıt hacmi sınırlıdır.",
        report_mode="compact",
    )
    prs = _deck(inputs)

    rails = {_text(slide, "rail_label") for slide in prs.slides}
    assert "Tematik kanıt sentezi" not in rails
    assert "Değerlendirme ve sonuç" not in rails
    assert ("3", "Belirsizlikler ve araştırma boşlukları") in [row[:2] for row in _contents(prs)]
    assert "Seyrek kanıt yalnız sınırlı bir sonuç destekler [S01]." in _deck_text(prs)
    for number, slide in enumerate(prs.slides, 1):
        if _shape(slide, "slide_number") is not None:
            assert _text(slide, "slide_number") == str(number)


def test_standard_report_renders_every_theme_with_its_number():
    inputs = _minimal_report_inputs()
    inputs["synthesis_package"] = _package(
        [
            SynthesisSection(
                title="Klinik Doğruluk",
                synthesis="Klinik doğruluk modellerde yüksektir [S01].",
                consensus="Modeller erken teşhiste etkilidir.",
                implications="Klinik iş akışını destekler.",
            ),
            SynthesisSection(title="İş Akışı Entegrasyonu", synthesis="Entegrasyon uyum gerektirir [S01]."),
        ]
    )
    prs = _deck(inputs)

    by_heading = {_text(slide, "content_heading"): _text(slide, "rail_number") for slide in prs.slides}
    assert by_heading["Klinik Doğruluk"] == "3.1"
    assert by_heading["İş Akışı Entegrasyonu"] == "3.2"
    assert "Değerlendirme ve sonuç" in {_text(slide, "rail_label") for slide in prs.slides}
    rows = [row[:2] for row in _contents(prs)]
    assert ("3.1", "Klinik Doğruluk") in rows
    assert ("3.2", "İş Akışı Entegrasyonu") in rows


def test_every_figure_slide_shows_its_own_image():
    """Run 01M2FYR67BS5WXFMEVY2RT1EHP shipped four figure slides; only the first image opened."""
    images = {name: _png(color) for name, color in (("A", "red"), ("B", "green"), ("C", "blue"))}
    inputs = _minimal_report_inputs()
    inputs["research_figures"] = [
        SimpleNamespace(
            title=f"Figür başlığı {name}",
            caption=f"Şekil {name}",
            attribution="S01",
            rights_statement="CC-BY 4.0",
            data=data,
            observation_hash=f"h{name}",
        )
        for name, data in images.items()
    ]
    # Included in the report, but no image of it was exported.
    inputs["figure_observations"] = [
        SimpleNamespace(
            image_hash="h-no-image",
            recommended_section="Görseli olmayan gözlem",
            selection_reason="Veri kümesi örneği",
            caption="Figure 2: Dataset examples.",
            source_label="S01",
            main_findings=["Sağ tarafta plevral efüzyon."],
            limitations=[],
        )
    ]
    prs = _deck(inputs)

    assert _broken_image_links(prs) == []
    assert "Görseli olmayan gözlem" not in _deck_text(prs)
    shown = {
        _text(slide, "content_heading"): _shape(slide, "figure.asset").image.blob
        for slide in prs.slides
        if _text(slide, "eyebrow").startswith("FİGÜR")
    }
    assert shown == {f"Figür başlığı {name}": data for name, data in images.items()}


def test_long_prose_continues_on_the_next_slide_without_being_cut_or_shrunk():
    inputs = _minimal_report_inputs()
    paragraphs = [" ".join(LONG_SENTENCES * 5), " ".join(LONG_SENTENCES * 4)]
    inputs["synthesis_package"] = _package(
        [SynthesisSection(title="Klinik Doğruluk ve Hassasiyet", synthesis="\n\n".join(paragraphs))]
    )
    prs = _deck(inputs)

    theme_slides = [
        slide
        for slide in prs.slides
        if _text(slide, "rail_number") == "3.1" and _text(slide, "content_heading").startswith("Klinik")
    ]
    assert len(theme_slides) > 1
    assert _text(theme_slides[0], "content_heading") == "Klinik Doğruluk ve Hassasiyet"
    assert all(_text(slide, "content_heading").endswith("(devam)") for slide in theme_slides[1:])
    shown = [
        shape.text_frame.text
        for slide in theme_slides
        for shape in slide.shapes
        if shape.name.startswith("sections[].synthesis")
    ]
    assert " ".join(shown) == " ".join(paragraphs)
    sizes = {run.font.size.pt for slide in prs.slides for shape in slide.shapes if shape.has_text_frame for paragraph in shape.text_frame.paragraphs for run in paragraph.runs}
    assert min(sizes) >= 12


def test_every_text_box_holds_its_measured_text_and_stays_in_its_region():
    prs = _deck(_long_inputs())

    for number, slide in enumerate(prs.slides, 1):
        for shape in slide.shapes:
            assert shape.left.pt + shape.width.pt <= 960.5, (number, shape.name)
            assert shape.top.pt + shape.height.pt <= 540.5, (number, shape.name)
            if not shape.has_text_frame or not shape.text_frame.text:
                continue
            paragraph = shape.text_frame.paragraphs[0]
            run = paragraph.runs[0]
            size = run.font.size.pt
            style = TextStyle(size, bold=bool(run.font.bold), leading=paragraph.line_spacing.pt / size)
            needed = text_height(shape.text_frame.text, style, shape.width.pt)
            assert needed <= shape.height.pt + 0.5, (number, shape.name, needed, shape.height.pt)
            assert run.font.name == "Arial", (number, shape.name)
            footer = shape.name in {"footer_note", "slide_number"}
            if number not in (1, len(prs.slides)) and not footer and shape.name != "logo":
                assert shape.top.pt + shape.height.pt <= FOOTER_TOP + 0.5, (number, shape.name)


def test_presentation_asserts_nothing_the_run_did_not_establish():
    """Empty sections used to become verdicts on the slides."""
    inputs = _minimal_report_inputs()
    inputs["coverage"] = {}
    inputs["synthesis_package"] = _package(
        [
            SynthesisSection(
                title="Tema bir",
                synthesis="Birinci taslak [S01].",
                reader_note="Bu tema tek bir metinde birleştirilemedi.",
            ),
            SynthesisSection(title="Tema iki", synthesis="İkinci tema [S01]."),
        ],
        cross_study_assessment="",
        conclusion="",
        uncertainty="",
        generated_by_llm=False,
        generation_status="partial",
    )
    deck = _deck_text(_deck(inputs))

    for verdict in (
        "genel tutarlılık saptanmıştır",
        "doğrulanmış kanıtlarla yanıtlanmıştır",
        "belirsizlik veya veri boşluğu bildirilmedi",
        "doğrulanmış kanıt kaydıyla desteklenmiştir",
        "bulguları desteklemektedir",
        "Tamamlandı",
        "100%",
        "%100",
    ):
        assert verdict not in deck, verdict
    assert "Bu tema tek bir metinde birleştirilemedi." in deck


def test_frame_formats_dates_and_contents_list_the_themes():
    inputs = _minimal_report_inputs()
    inputs["scope"] = {
        "start_date": "2023-09-15T08:37:12.273500Z",
        "end_date": "2026-09-14T10:00:00Z",
        "query_intent": "Transformer mimarisi",
    }
    inputs["synthesis_package"] = _package(
        [
            SynthesisSection(title="Dikkat Mekanizmaları", synthesis="Self-attention [S01]."),
            SynthesisSection(title="Pozisyonel Kodlama", synthesis="Sinüzoidal kodlama [S01]."),
        ]
    )
    prs = _deck(inputs)

    rows = [row[:2] for row in _contents(prs)]
    assert ("3.1", "Dikkat Mekanizmaları") in rows
    assert ("3.2", "Pozisyonel Kodlama") in rows
    frame = next(slide for slide in prs.slides if _text(slide, "content_heading") == "Ana soru ve kapsam")
    assert "15.09.2023 – 14.09.2026" in _texts(frame)
    assert not any("2023-09-15T" in text for text in _texts(frame))


def test_contents_page_numbers_point_at_their_slides():
    prs = _deck(_long_inputs())
    rows = _contents(prs)

    assert rows
    for marker, entry, page in rows:
        if not page:
            continue
        slide = prs.slides[int(page) - 1]
        if marker.startswith(("3.", "Ek")):
            assert _text(slide, "rail_number") == marker, (marker, page)
        if marker.startswith("3."):
            assert _text(slide, "content_heading") == entry


def test_findings_slides_only_for_themes_with_findings_and_appendices_summarise():
    inputs = _minimal_report_inputs()
    sources = [
        SimpleNamespace(
            id=f"src-{i}",
            title=f"Araştırma Kaynağı {i}",
            url=f"https://www.example.org/s{i}",
            connector_id="crossref",
            metadata_json={"research_scope_role": "primary_in_scope"},
        )
        for i in range(1, 7)
    ]
    claims = [
        SimpleNamespace(
            id=f"claim-{i}",
            text=f"Doğrulanmış iddia {i}",
            status="supported",
            confidence=0.9,
            audit={"question_relevance": 0.85},
        )
        for i in range(1, 6)
    ]
    inputs.update(
        {
            "sources": sources,
            "claims": claims,
            "reportable_claims": claims,
            "evidence_by_claim": {
                "claim-1": [(SimpleNamespace(), sources[0]), (SimpleNamespace(), sources[1])],
                "claim-2": [(SimpleNamespace(), sources[1])],
            },
            "figures": {
                "16a_research_contribution_landscape.png": _png("blue", (400, 300)),
                "16b_theme_evidence_map.png": _png("green", (400, 300)),
            },
            "synthesis_package": _package(
                [
                    SynthesisSection(
                        title="Tema 1",
                        synthesis="Tema 1 sentezi [S01].",
                        consensus="Tema 1 konsensus",
                        implications="Tema 1 çıkarım",
                    ),
                    SynthesisSection(title="Tema 2", synthesis="Tema 2 sentezi [S02]."),
                ]
            ),
        }
    )
    prs = _deck(inputs)

    findings = [slide for slide in prs.slides if _shape(slide, "sections[].consensus") is not None]
    assert [_text(slide, "rail_number") for slide in findings] == ["3.1"]

    landscape = next(slide for slide in prs.slides if _text(slide, "rail_number") == "Ek B")
    assert _shape(landscape, "contribution_landscape").shape_type == 13
    assert _shape(landscape, "theme_evidence_map").shape_type == 13

    catalog = next(slide for slide in prs.slides if _text(slide, "rail_number") == "Ek C")
    assert re.fullmatch(r"İddia sayısına göre ilk \d+ / 6 kaynak", _text(catalog, "catalog_tag"))
    # Sorted by how many claims a source supports: S02 backs two, S01 one.
    assert [_text(catalog, "catalog_1_1"), _text(catalog, "catalog_2_1")] == ["S02", "S01"]

    register = next(slide for slide in prs.slides if _text(slide, "rail_number") == "Ek D")
    assert _text(register, "claims_1_1") == "C01"
    assert _text(register, "claims_1_6") == "S01, S02"


def test_appendix_tables_name_where_the_full_list_is():
    prs = _deck(_long_inputs())

    catalog = next(slide for slide in prs.slides if _text(slide, "rail_number") == "Ek C")
    shown, total = re.fullmatch(r"İddia sayısına göre ilk (\d+) / (\d+) kaynak", _text(catalog, "catalog_tag")).groups()
    assert int(shown) < int(total) == 40
    assert _text(catalog, "footer_note") == "Tam liste: Word raporu Ek C ve 05_source_catalog.csv"
    # Site names are stripped from titles; the domain has its own column.
    assert "Springer Nature Link" not in _text(catalog, "catalog_1_2")
    assert _text(catalog, "catalog_1_3") == "onlinelibrary.wiley.com"

    register = next(slide for slide in prs.slides if _text(slide, "rail_number") == "Ek D")
    assert _text(register, "footer_note") == "Tam kayıt: Word raporu Ek D ve 04_claim_ledger.jsonl"


def test_figure_captions_do_not_repeat_their_prefixes():
    prs = _deck(_long_inputs())
    captions = [_text(slide, "figure.caption_attribution") for slide in prs.slides]
    caption = next(text for text in captions if text.startswith("Figure 2"))
    assert "Kaynak: Kaynak:" not in caption and "Telif: Telif:" not in caption
    heading = next(
        _text(slide, "content_heading") for slide in prs.slides if _text(slide, "eyebrow") == "FİGÜR 1 · S47"
    )
    assert heading == "Hacimsel BT görüntülerinin transvers kesitleri ve karşılık gelen raporlar"


def test_empty_summary_uses_the_synthesis_wording():
    inputs = _minimal_report_inputs()
    inputs["executive_summary"] = ""
    inputs["synthesis_package"] = _package([], executive_summary="", cross_study_assessment="", conclusion="", uncertainty="", generated_by_llm=False)
    prs = _deck(inputs)

    summary = next(slide for slide in prs.slides if _text(slide, "rail_label") == "Özet")
    text = "\n".join(_texts(summary))
    assert _REPORT_WITHOUT_SUMMARY[True] in text
    assert "eşiğ" not in text and "üretilemedi" not in text


def test_english_report_uses_english_labels():
    inputs = _minimal_report_inputs()
    inputs["language"] = "en"
    prs = _deck(inputs)
    deck = _deck_text(prs)

    assert "Contents" in deck and "Thank you" in deck
    assert "İçindekiler" not in deck and "Teşekkürler" not in deck


def test_base_template_embeds_no_fonts_and_names_arial():
    """Roboto embedded by the Google Slides export was not recognised by PowerPoint."""
    with zipfile.ZipFile(TEMPLATE_PATH) as package:
        names = package.namelist()
        assert not [name for name in names if name.startswith("ppt/fonts/")]
        assert "embeddedFontLst" not in package.read("ppt/presentation.xml").decode("utf-8")
        assert not [name for name in names if name.startswith("ppt/slides/")]
        for theme in (name for name in names if "/theme" in name):
            fonts = re.findall(r'<a:(?:major|minor)Font><a:latin typeface="([^"]+)"', package.read(theme).decode("utf-8"))
            assert set(fonts) <= {"Arial"}
    assert LOGO_WHITE.exists() and LOGO_COLOR.exists()
