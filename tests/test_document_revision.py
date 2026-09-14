from __future__ import annotations

import hashlib
import io
import json
import zipfile

import pytest
from conftest import acting_principal, api_headers, ensure_test_user
from docx import Document
from fake_redis import FakeRedis
from fastapi.testclient import TestClient
from pptx import Presentation
from pptx.util import Cm

from research_platform.api import app
from research_platform.config import get_settings
from research_platform.db import SessionLocal, create_schema
from research_platform.document_revision import (
    DOCX_MEDIA_TYPE,
    PPTX_MEDIA_TYPE,
    DocumentRevisionService,
    RevisionValidationError,
    _apply_format_overrides,
    apply_revision_plan,
    editable_targets,
    office_format_context,
    report_model_from_manifest,
)
from research_platform.queueing import revision_job_id_for
from research_platform.repository import Repository, RevisionConflict
from research_platform.schemas import ResearchProtocol, RevisionPlan, RevisionStatus, RunStatus
from research_platform.storage import ObjectStore
from research_platform.worker import _recover_interrupted_jobs


def _manifest() -> dict:
    protocol = ResearchProtocol(
        title="Versioned report",
        primary_question="Which evidence supports this versioned research report?",
        report_language="en",
        budget={"max_wall_minutes": 30},
    )
    return {
        "protocol": protocol.model_dump(mode="json"),
        "coverage": {"sufficient": True},
        "synthesis": {
            "executive_summary": "The retained evidence supports the finding [S01].",
            "sections": [
                {
                    "title": "Evidence",
                    "synthesis": "The primary study reports the finding [S01].",
                    "consensus": "One source is represented [S01].",
                    "disagreements": "",
                    "implications": "The result should be interpreted cautiously [S01].",
                    "source_ids": ["source-1"],
                    "claim_ids": ["claim-1"],
                }
            ],
            "cross_study_assessment": "Only one source was available [S01].",
            "conclusion": "The conclusion remains bounded by the evidence [S01].",
            "uncertainty": "Independent replication is still needed [S01].",
            "study_profiles": [],
            "generated_by_llm": False,
        },
    }


def _plan(**operation) -> RevisionPlan:
    return RevisionPlan.model_validate(
        {
            "summary": "Apply the requested report edit.",
            "operations": [
                {
                    "operation": "set_title",
                    "target": "document.title",
                    "instruction": "Use a more direct title.",
                    "replacement": "A clearer versioned report",
                    **operation,
                }
            ],
        }
    )


def _bytes_for_office_files() -> tuple[bytes, bytes]:
    docx = io.BytesIO()
    Document().save(docx)
    pptx = io.BytesIO()
    Presentation().save(pptx)
    return docx.getvalue(), pptx.getvalue()


def _bundle(files: dict[str, bytes]) -> bytes:
    output = io.BytesIO()
    with zipfile.ZipFile(output, "w", zipfile.ZIP_DEFLATED) as archive:
        for name, data in files.items():
            archive.writestr(name, data)
    return output.getvalue()


async def _revision_fixture(repo: Repository):
    model = report_model_from_manifest(_manifest())
    run = await repo.create_run(ResearchProtocol.model_validate(_manifest()["protocol"]))
    await repo.update_run(run.id, status=RunStatus.COMPLETED.value)
    initial = await repo.bootstrap_document_revision(
        run.id,
        report_model=model,
        artifact_versions=[
            {
                "logical_name": "report.docx",
                "media_type": DOCX_MEDIA_TYPE,
                "object_key": "old.docx",
                "sha256": "a" * 64,
                "size_bytes": 10,
            },
            {
                "logical_name": "report.pptx",
                "media_type": PPTX_MEDIA_TYPE,
                "object_key": "old.pptx",
                "sha256": "b" * 64,
                "size_bytes": 20,
            },
        ],
        citations=[],
        requested_by=str(run.owner_id),
    )
    draft = await repo.create_document_revision(
        run.id,
        target_artifact_name="report.pptx",
        feedback="Use a clearer title",
        base_revision_id=initial.id,
        parent_revision_id=None,
        idempotency_key=None,
        channel="api",
        conversation_id=None,
    )
    return run, initial, draft


def test_manifest_becomes_a_stable_addressable_report_model():
    first = report_model_from_manifest(_manifest())
    second = report_model_from_manifest(_manifest())

    assert first == second
    targets = editable_targets(first)
    section_id = first["synthesis"]["sections"][0]["id"]
    assert f"section:{section_id}.synthesis" in targets
    assert targets["document.title"] == "Versioned report"


def test_a_revision_changes_only_approved_targets_and_keeps_source_labels():
    original = report_model_from_manifest(_manifest())
    updated, overrides = apply_revision_plan(original, _plan())

    assert updated["title"] == "A clearer versioned report"
    assert original["title"] == "Versioned report"
    assert updated["synthesis"] == original["synthesis"]
    assert overrides == {}


def test_content_operations_must_update_both_office_formats():
    model = report_model_from_manifest(_manifest())
    with pytest.raises(RevisionValidationError, match="shared DOCX/PPTX"):
        apply_revision_plan(model, _plan(format_scope="pptx"))


def test_a_revision_cannot_remove_or_invent_citations():
    model = report_model_from_manifest(_manifest())
    for replacement in (
        "The retained evidence supports the finding.",
        "The retained evidence supports the finding [S99].",
    ):
        plan = RevisionPlan.model_validate(
            {
                "summary": "Rewrite the summary.",
                "operations": [
                    {
                        "operation": "rewrite_section",
                        "target": "executive_summary",
                        "instruction": "Rewrite it.",
                        "replacement": replacement,
                    }
                ],
            }
        )
        with pytest.raises(RevisionValidationError):
            apply_revision_plan(model, plan)


def test_numeric_findings_and_quotations_remain_locked():
    model = report_model_from_manifest(_manifest())
    model["synthesis"]["executive_summary"] = (
        'The response was 18.5% and the authors wrote “clinically meaningful” [S01].'
    )
    for replacement in (
        'The response was 21% and the authors wrote “clinically meaningful” [S01].',
        'The response was 18.5% and the authors wrote “uncertain” [S01].',
    ):
        plan = RevisionPlan.model_validate(
            {
                "summary": "Rewrite the summary.",
                "operations": [
                    {
                        "operation": "rewrite_section",
                        "target": "executive_summary",
                        "instruction": "Rewrite it.",
                        "replacement": replacement,
                    }
                ],
            }
        )
        with pytest.raises(RevisionValidationError, match="locked"):
            apply_revision_plan(model, plan)


def test_format_overrides_are_bounded_and_applied_without_editing_ooxml():
    document = Document()
    document.add_paragraph("Heading", style="Heading 1")
    docx = io.BytesIO()
    document.save(docx)
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    shape = slide.shapes.add_textbox(Cm(1), Cm(1), Cm(5), Cm(2))
    shape.name = "Summary box"
    shape.text = "Summary"
    pptx = io.BytesIO()
    presentation.save(pptx)

    updated_docx, updated_pptx = _apply_format_overrides(
        docx.getvalue(),
        pptx.getvalue(),
        {
            "docx.style:Heading 1.font_size_pt": 22,
            "pptx.slide:1.shape:Summary box.bounds": {"width_cm": 8, "height_cm": 3},
            "pptx.slide:1.shape:Summary box.font_size_pt": 18,
        },
    )

    assert Document(io.BytesIO(updated_docx)).styles["Heading 1"].font.size.pt == 22
    updated_shape = Presentation(io.BytesIO(updated_pptx)).slides[0].shapes[0]
    assert updated_shape.width == Cm(8)
    assert updated_shape.height == Cm(3)
    assert updated_shape.text_frame.paragraphs[0].runs[0].font.size.pt == 18

    with pytest.raises(RevisionValidationError, match="outside"):
        _apply_format_overrides(
            docx.getvalue(),
            pptx.getvalue(),
            {"pptx.slide:1.shape:Summary box.bounds": {"width_cm": 200}},
        )
    with pytest.raises(RevisionValidationError, match="exceeds slide bounds"):
        _apply_format_overrides(
            docx.getvalue(),
            pptx.getvalue(),
            {
                "pptx.slide:1.shape:Summary box.bounds": {
                    "left_cm": 30,
                    "width_cm": 8,
                }
            },
        )


def test_format_context_exposes_safe_shape_and_style_targets():
    document = Document()
    document.add_paragraph("Heading", style="Heading 1")
    docx = io.BytesIO()
    document.save(docx)
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    shape = slide.shapes.add_textbox(Cm(1), Cm(2), Cm(5), Cm(3))
    shape.name = "Summary box"
    shape.text = "Evidence summary"
    pptx = io.BytesIO()
    presentation.save(pptx)

    pptx_context = office_format_context(pptx.getvalue(), PPTX_MEDIA_TYPE)
    docx_context = office_format_context(docx.getvalue(), DOCX_MEDIA_TYPE)

    target = pptx_context["pptx"]["slides"][0]["shapes"][0]
    assert target["name"] == "Summary box"
    assert target["text"] == "Evidence summary"
    assert target["bounds_cm"] == {
        "left": 1.0,
        "top": 2.0,
        "width": 5.0,
        "height": 3.0,
    }
    assert "Heading 1" in docx_context["docx"]["styles"]


@pytest.mark.asyncio
async def test_planning_stops_for_explicit_user_approval_before_rendering():
    class PlanningLLM:
        async def complete_json(self, system: str, user: str) -> dict:
            assert "conservative research-report editor" in system
            assert "Use a clearer title" in user
            return _plan().model_dump(mode="json")

        def drain_metrics(self) -> list[dict]:
            return [{"model": "test-planner"}]

    await create_schema()
    async with SessionLocal() as session:
        repo = Repository(session, actor=acting_principal())
        _run, _initial, draft = await _revision_fixture(repo)
        planned = await DocumentRevisionService(
            repo, ObjectStore(get_settings()), llm=PlanningLLM(), settings=get_settings()
        ).plan(draft.id)

        assert planned.status == RevisionStatus.AWAITING_PLAN_APPROVAL.value
        assert planned.report_model == {}
        assert planned.edit_plan["operations"][0]["target"] == "document.title"


@pytest.mark.asyncio
async def test_cancelling_during_planning_is_not_overwritten_by_the_llm_result():
    await create_schema()
    async with SessionLocal() as session:
        repo = Repository(session, actor=acting_principal())
        run, _initial, draft = await _revision_fixture(repo)

        class CancellingLLM:
            async def complete_json(self, _system: str, _user: str) -> dict:
                async with SessionLocal() as control_session:
                    await Repository(
                        control_session, actor=acting_principal()
                    ).update_document_revision(
                        run.id, draft.id, status=RevisionStatus.CANCELLED.value
                    )
                return _plan().model_dump(mode="json")

            def drain_metrics(self) -> list[dict]:
                return []

        planned = await DocumentRevisionService(
            repo, ObjectStore(get_settings()), llm=CancellingLLM(), settings=get_settings()
        ).plan(draft.id)

        assert planned.status == RevisionStatus.CANCELLED.value
        assert planned.edit_plan is None


@pytest.mark.asyncio
async def test_cancelled_revision_does_not_accept_late_render_outputs():
    await create_schema()
    async with SessionLocal() as session:
        repo = Repository(session, actor=acting_principal())
        run, _initial, draft = await _revision_fixture(repo)
        await repo.update_document_revision(
            run.id, draft.id, status=RevisionStatus.VALIDATING.value
        )
        await repo.update_document_revision(
            run.id, draft.id, status=RevisionStatus.CANCELLED.value
        )

        result = await repo.save_revision_outputs(
            run.id,
            draft.id,
            report_model=report_model_from_manifest(_manifest()),
            format_overrides={},
            validation={"passed": True},
            artifact_versions=[
                {
                    "logical_name": "late.pptx",
                    "parent_version_id": None,
                    "media_type": PPTX_MEDIA_TYPE,
                    "object_key": "late.pptx",
                    "sha256": "c" * 64,
                    "size_bytes": 30,
                }
            ],
            citations=[],
        )

        assert result.status == RevisionStatus.CANCELLED.value
        assert await repo.list_artifact_versions(run.id, revision_id=draft.id) == []


@pytest.mark.asyncio
async def test_worker_restart_recovers_an_interrupted_revision_job():
    await create_schema()
    async with SessionLocal() as session:
        repo = Repository(session, actor=acting_principal())
        _run, _initial, draft = await _revision_fixture(repo)
        await repo.update_document_revision(
            draft.run_id, draft.id, status=RevisionStatus.PLANNING.value
        )

    redis = FakeRedis()
    await _recover_interrupted_jobs({"redis": redis})

    async with SessionLocal() as session:
        recovered = await Repository(
            session, actor=acting_principal()
        ).get_document_revision(draft.run_id, draft.id)
    assert recovered.status == RevisionStatus.QUEUED.value
    assert revision_job_id_for(draft.id) in redis.queue


@pytest.mark.asyncio
async def test_rendering_regenerates_both_office_files_and_their_bundles():
    await create_schema()
    docx, pptx = _bytes_for_office_files()
    manifest = json.dumps(_manifest()).encode()
    initial_files = {"report.docx": docx, "report.pptx": pptx}
    result_bundle = _bundle(initial_files)
    research_bundle = _bundle(
        {**initial_files, "10_reproducibility_manifest.json": manifest}
    )
    store = ObjectStore(get_settings())

    async with SessionLocal() as session:
        repo = Repository(session, actor=acting_principal())
        run = await repo.create_run(
            ResearchProtocol.model_validate(_manifest()["protocol"])
        )
        await repo.update_run(run.id, status=RunStatus.COMPLETED.value)
        payloads = {
            "report.docx": (DOCX_MEDIA_TYPE, docx),
            "report.pptx": (PPTX_MEDIA_TYPE, pptx),
            "result_bundle.zip": ("application/zip", result_bundle),
            "research_bundle.zip": ("application/zip", research_bundle),
            "10_reproducibility_manifest.json": ("application/json", manifest),
        }
        for name, (media_type, data) in payloads.items():
            key = f"runs/{run.id}/{name}"
            await store.put(key, data, media_type)
            await repo.save_artifact(run.id, name, media_type, key, len(data))

        service = DocumentRevisionService(repo, store, settings=get_settings())
        initial = await service.ensure_initial_revision(run.id)
        draft = await repo.create_document_revision(
            run.id,
            target_artifact_name="report.pptx",
            feedback="Use a clearer title",
            base_revision_id=initial.id,
            parent_revision_id=None,
            idempotency_key=None,
            channel="api",
            conversation_id=None,
        )
        await repo.update_document_revision(
            run.id, draft.id, edit_plan=_plan().model_dump(mode="json")
        )

        rendered = await service.render(draft.id)
        versions = await repo.list_artifact_versions(run.id, revision_id=draft.id)

        assert rendered.status == "awaiting_revision_approval"
        assert rendered.report_model["title"] == "A clearer versioned report"
        assert rendered.validation["passed"] is True
        by_name = {item.logical_name: await store.get(item.object_key) for item in versions}
        Document(io.BytesIO(by_name["report.docx"]))
        Presentation(io.BytesIO(by_name["report.pptx"]))
        for bundle_name in ("result_bundle.zip", "research_bundle.zip"):
            with zipfile.ZipFile(io.BytesIO(by_name[bundle_name])) as archive:
                assert archive.read("report.docx") == by_name["report.docx"]
                assert archive.read("report.pptx") == by_name["report.pptx"]


@pytest.mark.asyncio
async def test_accepting_a_revision_atomically_moves_the_compatibility_pointers():
    await create_schema()
    model = report_model_from_manifest(_manifest())
    async with SessionLocal() as session:
        repo = Repository(session, actor=acting_principal())
        run = await repo.create_run(
            ResearchProtocol.model_validate(_manifest()["protocol"])
        )
        await repo.update_run(run.id, status=RunStatus.COMPLETED.value)
        await repo.save_artifact(run.id, "report.docx", DOCX_MEDIA_TYPE, "old.docx", 10)
        await repo.save_artifact(run.id, "report.pptx", PPTX_MEDIA_TYPE, "old.pptx", 20)
        initial = await repo.bootstrap_document_revision(
            run.id,
            report_model=model,
            artifact_versions=[
                {
                    "logical_name": "report.docx",
                    "media_type": DOCX_MEDIA_TYPE,
                    "object_key": "old.docx",
                    "sha256": "a" * 64,
                    "size_bytes": 10,
                },
                {
                    "logical_name": "report.pptx",
                    "media_type": PPTX_MEDIA_TYPE,
                    "object_key": "old.pptx",
                    "sha256": "b" * 64,
                    "size_bytes": 20,
                },
            ],
            citations=[],
            requested_by=str(run.owner_id),
        )
        draft = await repo.create_document_revision(
            run.id,
            target_artifact_name="report.pptx",
            feedback="Use a clearer title",
            base_revision_id=initial.id,
            parent_revision_id=None,
            idempotency_key="accept-pointer-test",
            channel="api",
            conversation_id=None,
        )
        revised_model, _ = apply_revision_plan(model, _plan())
        await repo.update_document_revision(
            run.id, draft.id, status=RevisionStatus.VALIDATING.value
        )
        await repo.save_revision_outputs(
            run.id,
            draft.id,
            report_model=revised_model,
            format_overrides={},
            validation={"passed": True},
            artifact_versions=[
                {
                    "logical_name": "report.docx",
                    "parent_version_id": None,
                    "media_type": DOCX_MEDIA_TYPE,
                    "object_key": "new.docx",
                    "sha256": "c" * 64,
                    "size_bytes": 30,
                },
                {
                    "logical_name": "report.pptx",
                    "parent_version_id": None,
                    "media_type": PPTX_MEDIA_TYPE,
                    "object_key": "new.pptx",
                    "sha256": "d" * 64,
                    "size_bytes": 40,
                },
            ],
            citations=[],
        )

        accepted = await repo.accept_document_revision(run.id, draft.id)
        artifacts = {row.name: row for row in await repo.list_artifacts(run.id)}

        assert accepted.status == "accepted"
        assert (await repo.get_document_revision(run.id, initial.id)).status == "superseded"
        assert artifacts["report.docx"].object_key == "new.docx"
        assert artifacts["report.pptx"].object_key == "new.pptx"
        with pytest.raises(RevisionConflict, match="stale"):
            await repo.create_document_revision(
                run.id,
                target_artifact_name="report.pptx",
                feedback="Edit the old base",
                base_revision_id=initial.id,
                parent_revision_id=None,
                idempotency_key=None,
                channel="api",
                conversation_id=None,
            )
        restored = await repo.restore_document_revision(
            run.id, initial.id, channel="panel"
        )
        restored_versions = await repo.list_artifact_versions(
            run.id, revision_id=restored.id
        )
        assert restored.status == "awaiting_revision_approval"
        assert {item.object_key for item in restored_versions} == {
            "old.docx",
            "old.pptx",
        }
        await repo.accept_document_revision(run.id, restored.id)
        restored_artifacts = {
            row.name: row for row in await repo.list_artifacts(run.id)
        }
        assert restored_artifacts["report.docx"].object_key == "old.docx"
        assert restored_artifacts["report.pptx"].object_key == "old.pptx"


@pytest.mark.asyncio
async def test_api_lazily_versions_existing_artifacts_and_keeps_download_metadata():
    await ensure_test_user()
    docx, pptx = _bytes_for_office_files()
    manifest = json.dumps(_manifest()).encode()
    names = {"report.docx": docx, "report.pptx": pptx}
    result_bundle = _bundle(names)
    research_bundle = _bundle({**names, "10_reproducibility_manifest.json": manifest})
    store = ObjectStore(get_settings())

    async with SessionLocal() as session:
        repo = Repository(session, actor=acting_principal())
        run = await repo.create_run(
            ResearchProtocol.model_validate(_manifest()["protocol"])
        )
        await repo.update_run(run.id, status=RunStatus.COMPLETED.value)
        payloads = {
            "report.docx": (DOCX_MEDIA_TYPE, docx),
            "report.pptx": (PPTX_MEDIA_TYPE, pptx),
            "result_bundle.zip": ("application/zip", result_bundle),
            "research_bundle.zip": ("application/zip", research_bundle),
            "10_reproducibility_manifest.json": ("application/json", manifest),
        }
        for name, (media_type, data) in payloads.items():
            key = f"runs/{run.id}/{name}"
            await store.put(key, data, media_type)
            await repo.save_artifact(run.id, name, media_type, key, len(data))
        run_id = run.id

    with TestClient(app) as client:
        client.headers.update(api_headers())
        versions_response = client.get(f"/v1/research-runs/{run_id}/artifact-versions")
        assert versions_response.status_code == 200, versions_response.text
        versions = versions_response.json()
        assert {item["logical_name"] for item in versions} == {
            "report.docx",
            "report.pptx",
            "result_bundle.zip",
            "research_bundle.zip",
        }
        assert all(item["revision_number"] == 1 for item in versions)

        pptx_version = next(item for item in versions if item["logical_name"] == "report.pptx")
        downloaded = client.get(f"/v1/artifact-versions/{pptx_version['id']}")
        assert downloaded.status_code == 200
        assert downloaded.content == pptx
        assert downloaded.headers["x-revision-id"] == pptx_version["revision_id"]
        assert downloaded.headers["x-artifact-version-id"] == pptx_version["id"]

        created = client.post(
            f"/v1/research-runs/{run_id}/artifacts/report.pptx/revisions",
            json={"channel": "api"},
            headers={"Idempotency-Key": hashlib.sha256(run_id.encode()).hexdigest()},
        )
        assert created.status_code == 200, created.text
        assert created.json()["status"] == "awaiting_feedback"
        queued = client.post(
            f"/v1/research-runs/{run_id}/revisions/{created.json()['id']}/feedback",
            json={"feedback": "Use a more direct title."},
        )
        assert queued.status_code == 200, queued.text
        assert queued.json()["status"] == "queued"
