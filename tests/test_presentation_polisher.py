"""Tests for the presentation polisher: the agent has full authority, delivery never fails."""

from __future__ import annotations

import asyncio
import io
import json
import sys
from datetime import datetime, timedelta, timezone
from pathlib import Path
from unittest.mock import AsyncMock, patch

import httpx
import pptx
import pytest
from pptx.util import Inches

from research_platform.config import Settings
from research_platform.diagnostics import event_severity
from research_platform.presentation_polisher import (
    POLISH_EVENT,
    PolishResult,
    polish_and_record,
    polish_presentation,
)
from scripts import presentation_polisher_outline as outline_helper
from scripts import presentation_polisher_render as render_helper
from scripts import presentation_polisher_service as service

PPTX_MEDIA_TYPE = "application/vnd.openxmlformats-officedocument.presentationml.presentation"
TOKEN = "test-polisher-token"


def _deck(text: str = "Özgün başlık [S01]") -> bytes:
    prs = pptx.Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    box = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(6), Inches(1))
    box.text_frame.text = text
    buffer = io.BytesIO()
    prs.save(buffer)
    return buffer.getvalue()


def _settings(**overrides) -> Settings:
    values = {
        "presentation_polisher_enabled": True,
        "presentation_polisher_url": "http://127.0.0.1:3942",
        "presentation_polisher_timeout_s": 2.0,
        "presentation_polisher_token": TOKEN,
    }
    values.update(overrides)
    return Settings(**values)


def _reply(status: str, reason: str, content: bytes, **headers: str) -> httpx.Response:
    return httpx.Response(
        200,
        content=content,
        headers={
            "Content-Type": PPTX_MEDIA_TYPE,
            "X-Presentation-Polisher-Status": status,
            "X-Presentation-Polisher-Reason": reason,
            **headers,
        },
    )


class RecordingRepo:
    def __init__(self) -> None:
        self.events: list[tuple[str, str, dict]] = []

    async def event(self, run_id: str, event_type: str, payload: dict | None = None) -> int:
        self.events.append((run_id, event_type, payload or {}))
        return len(self.events)


# --------------------------------------------------------------------------- client


@pytest.mark.asyncio
async def test_disabled_polisher_is_skipped_without_network():
    with patch("httpx.AsyncClient.post", new_callable=AsyncMock) as post:
        result = await polish_presentation(
            b"original", settings=Settings(presentation_polisher_enabled=False)
        )
    assert result == PolishResult(data=b"original", status="skipped", reason="disabled")
    post.assert_not_called()


@pytest.mark.asyncio
async def test_shared_service_token_is_not_a_fallback():
    settings = _settings(presentation_polisher_token="", service_token="shared-token")
    with patch("httpx.AsyncClient.post", new_callable=AsyncMock) as post:
        result = await polish_presentation(b"original", settings=settings)
    assert (result.data, result.status, result.reason) == (b"original", "skipped", "no-token")
    post.assert_not_called()


@pytest.mark.asyncio
async def test_unreachable_and_slow_services_keep_the_original():
    data = _deck()
    unreachable = await polish_presentation(
        data, settings=_settings(presentation_polisher_url="http://127.0.0.1:59999")
    )
    assert unreachable.data == data
    assert unreachable.status == "failed"
    assert unreachable.reason.startswith("unreachable:")

    with patch("httpx.AsyncClient.post", new_callable=AsyncMock) as post:
        post.side_effect = httpx.ReadTimeout("slow agent")
        slow = await polish_presentation(data, settings=_settings())
    assert (slow.data, slow.status, slow.reason) == (data, "failed", "timeout")


@pytest.mark.asyncio
async def test_polished_deck_and_its_details_are_returned():
    original, polished = _deck(), _deck("Ajanın sunumu")
    with patch("httpx.AsyncClient.post", new_callable=AsyncMock) as post:
        post.return_value = _reply(
            "polished",
            "agy-SUCCESS",
            polished,
            **{
                "X-Presentation-Polisher-Agy-Status": "SUCCESS",
                "X-Presentation-Polisher-Duration-S": "412.5",
                "X-Presentation-Polisher-Turns": "31",
            },
        )
        result = await polish_presentation(original, run_id="run-1", settings=_settings())

    assert result == PolishResult(
        data=polished,
        status="polished",
        reason="agy-SUCCESS",
        agy_status="SUCCESS",
        duration_s=412.5,
        turns=31,
    )
    assert post.await_args.kwargs["headers"]["Authorization"] == f"Bearer {TOKEN}"
    assert post.await_args.kwargs["params"] == {"language": "tr", "run_id": "run-1"}


@pytest.mark.asyncio
async def test_service_decision_to_keep_the_original_is_passed_on():
    original = _deck()
    with patch("httpx.AsyncClient.post", new_callable=AsyncMock) as post:
        post.return_value = _reply("unchanged", "timeout-no-output", original)
        result = await polish_presentation(original, settings=_settings())
    assert (result.data, result.status, result.reason) == (original, "unchanged", "timeout-no-output")


@pytest.mark.asyncio
@pytest.mark.parametrize(
    ("response", "reason"),
    [
        (httpx.Response(503, json={"detail": "agy and LibreOffice are required"}), "http-503"),
        (
            httpx.Response(
                200,
                content=b"ok",
                headers={"Content-Type": "text/html", "X-Presentation-Polisher-Status": "polished"},
            ),
            "invalid-response",
        ),
        (_reply("polished", "agy-SUCCESS", b"not-pptx"), "invalid-response"),
        (httpx.Response(200, content=b"x", headers={"Content-Type": PPTX_MEDIA_TYPE}), "invalid-response"),
    ],
)
async def test_untrusted_responses_keep_the_original(response, reason):
    original = _deck()
    with patch("httpx.AsyncClient.post", new_callable=AsyncMock) as post:
        post.return_value = response
        result = await polish_presentation(original, settings=_settings())
    assert (result.data, result.status, result.reason) == (original, "failed", reason)


@pytest.mark.asyncio
async def test_every_attempt_is_recorded_on_the_run():
    repo = RecordingRepo()
    original, polished = _deck(), _deck("Ajanın sunumu")
    with patch("httpx.AsyncClient.post", new_callable=AsyncMock) as post:
        post.return_value = _reply("polished", "timeout-partial", polished)
        delivered = await polish_and_record(
            repo, "run-1", original, language="tr", settings=_settings(), stage="export"
        )
        post.side_effect = httpx.ConnectError("refused")
        kept = await polish_and_record(
            repo,
            "run-1",
            original,
            language="tr",
            settings=_settings(),
            stage="revision",
            revision_id="rev-2",
        )

    assert delivered == polished
    assert kept == original
    assert [(run, kind) for run, kind, _ in repo.events] == [("run-1", POLISH_EVENT)] * 2
    first, second = (payload for _, _, payload in repo.events)
    assert first["stage"] == "export"
    assert (first["status"], first["reason"]) == ("polished", "timeout-partial")
    assert (first["original_bytes"], first["output_bytes"]) == (len(original), len(polished))
    assert "revision_id" not in first
    assert (second["stage"], second["status"], second["revision_id"]) == ("revision", "failed", "rev-2")
    assert second["reason"] == "unreachable:ConnectError"


@pytest.mark.asyncio
async def test_nothing_is_recorded_while_the_feature_is_off():
    repo = RecordingRepo()
    delivered = await polish_and_record(
        repo,
        "run-1",
        b"deck",
        language="tr",
        settings=Settings(presentation_polisher_enabled=False),
        stage="export",
    )
    assert delivered == b"deck"
    assert repo.events == []


@pytest.mark.parametrize(
    ("status", "reason", "severity"),
    [
        ("polished", "agy-SUCCESS", "info"),
        ("unchanged", "agent-made-no-changes", "info"),
        ("unchanged", "timeout-no-output", "warning"),
        ("failed", "unreachable:ConnectError", "warning"),
        ("skipped", "no-token", "warning"),
    ],
)
def test_a_deck_that_was_not_polished_is_a_warning(status, reason, severity):
    assert event_severity(POLISH_EVENT, {"status": status, "reason": reason}) == severity


# --------------------------------------------------------------------------- service


def test_agent_environment_withholds_secrets_and_puts_the_service_python_first(monkeypatch):
    monkeypatch.setenv("SERVICE_TOKEN", "secret")
    monkeypatch.setenv("PRESENTATION_POLISHER_TOKEN", "secret")
    monkeypatch.setenv("DATABASE_URL", "secret")
    monkeypatch.setenv("PATH", "safe-path")
    environment = service._agent_environment("C:/LibreOffice/soffice.exe")
    path_key = next(name for name in environment if name.upper() == "PATH")
    assert environment[path_key].split(service.os.pathsep) == [
        str(Path(sys.executable).parent),
        "safe-path",
    ]
    assert environment["PRESENTATION_POLISHER_SOFFICE"] == "C:/LibreOffice/soffice.exe"
    assert (environment["PYTHONUTF8"], environment["PYTHONIOENCODING"]) == ("1", "utf-8")
    assert "SERVICE_TOKEN" not in environment
    assert "PRESENTATION_POLISHER_TOKEN" not in environment
    assert "DATABASE_URL" not in environment


@pytest.mark.parametrize("sandbox", [True, False])
def test_agy_gets_full_authority_and_the_prompt_last(monkeypatch, sandbox):
    monkeypatch.setattr(service, "AGY_SANDBOX", sandbox)
    monkeypatch.setattr(service, "AGY_MODEL", "gemini-test")
    command = service.agy_command("agy", "istem", 600.0)
    assert command[0] == "agy"
    assert "--dangerously-skip-permissions" in command
    assert command[command.index("--output-format") + 1] == "json"
    assert command[command.index("--print-timeout") + 1] == "600s"
    assert command[command.index("--model") + 1] == "gemini-test"
    assert ("--sandbox" in command) is sandbox
    assert "--mode" not in command
    assert command[-2:] == ["-p", "istem"]


def test_prompt_is_read_on_every_request_and_carries_the_contract(monkeypatch, tmp_path):
    prompt_file = tmp_path / "istem.md"
    prompt_file.write_text("Sunumu daha etkileyici yap.", encoding="utf-8")
    monkeypatch.setattr(service, "PROMPT_FILE", prompt_file)

    istanbul = timezone(timedelta(hours=3))
    prompt = service.build_prompt(
        "tr", 600.0, now=datetime(2026, 9, 16, 13, 50, tzinfo=istanbul)
    )
    assert prompt.startswith("Sunumu daha etkileyici yap.")
    for expected in (
        "sunum.pptx",
        "cilali.pptx",
        "cilali.tmp.pptx",
        "python render.py",
        "python outline.py cilali.pptx --check",
        "10 dakika",
        "Başlangıç saati 13:50",
        "en geç 13:58",
        "Türkçe (tr)",
    ):
        assert expected in prompt
    assert "[S" not in prompt  # citations are the agent's call, not a rule

    prompt_file.write_text("Yeni istem.", encoding="utf-8")
    assert service.build_prompt("en", 300.0).startswith("Yeni istem.")
    assert "English (en)" in service.build_prompt("en", 300.0)

    monkeypatch.setattr(service, "_MAX_PROMPT_CHARS", 50)
    with pytest.raises(service.PromptError, match="prompt-too-long"):
        service.build_prompt("tr", 600.0)

    monkeypatch.setattr(service, "PROMPT_FILE", tmp_path / "missing.md")
    with pytest.raises(service.PromptError, match="prompt-missing"):
        service.build_prompt("tr", 600.0)


def test_word_preferences_follow_the_report_language(monkeypatch, tmp_path):
    prompt_file = tmp_path / "istem.md"
    prompt_file.write_text("Sunumu toparla.", encoding="utf-8")
    terms_file = tmp_path / "kelimeler.json"
    terms_file.write_text(
        json.dumps(
            {"_açıklama": "not", "tr": {"figür": "şekil"}, "en": {}}, ensure_ascii=False
        ),
        encoding="utf-8",
    )
    monkeypatch.setattr(service, "PROMPT_FILE", prompt_file)
    monkeypatch.setattr(service, "TERMS_FILE", terms_file)

    turkish = service.build_prompt("tr", 600.0)
    # After the user's prompt and before the fixed rules, as readable UTF-8 JSON.
    assert turkish.startswith("Sunumu toparla.")
    rules = turkish.index("KELİME TERCİHLERİ (servis `kelimeler.json`")
    assert rules < turkish.index("TEKNİK ÇALIŞMA KURALLARI")
    assert '"figür": "şekil"' in turkish
    assert "_açıklama" not in turkish
    assert "KELİME TERCİHLERİ" not in service.build_prompt("en", 600.0)

    # Read on every request, like the prompt.
    terms_file.write_text(
        json.dumps({"tr": {"figür": "şekil", "veriseti": "veri kümesi"}}, ensure_ascii=False),
        encoding="utf-8",
    )
    assert '"veriseti": "veri kümesi"' in service.build_prompt("tr", 600.0)

    monkeypatch.setattr(service, "TERMS_FILE", tmp_path / "yok.json")
    assert "KELİME TERCİHLERİ" not in service.build_prompt("tr", 600.0)


@pytest.mark.parametrize(
    "content",
    [
        '{"tr": {"figür": "şekil",}}',
        '["figür", "şekil"]',
        '{"tr": ["figür", "şekil"]}',
        '{"tr": {"figür": ""}}',
        '{"tr": {" ": "şekil"}}',
        '{"en": {"figure": 3}}',
    ],
    ids=["syntax", "list", "language-list", "empty-use", "empty-avoid", "other-language"],
)
def test_a_broken_word_list_is_an_error_not_an_empty_list(monkeypatch, tmp_path, content):
    """Whoever edited the file must learn that the preference is not being applied."""
    terms_file = tmp_path / "kelimeler.json"
    terms_file.write_text(content, encoding="utf-8")
    monkeypatch.setattr(service, "TERMS_FILE", terms_file)
    with pytest.raises(service.PromptError, match="terms-invalid"):
        service.load_terms("tr")


def test_shipped_configuration_prefers_sekil_and_does_not_teach_figur():
    """The agent copies the prompt's own wording, so the prompt must not say "figür"."""
    assert service.load_terms("tr") == {"figür": "şekil"}
    assert service.load_terms("en") == {}
    assert "figür" not in service.PROMPT_FILE.read_text(encoding="utf-8").lower()


def test_envelope_survives_a_log_line_before_it():
    text = 'warming up\n{"status": "SUCCESS", "num_turns": 4, "response": "ok"}\n'
    assert service._parse_envelope(text)["num_turns"] == 4
    assert service._parse_envelope("not json") == {}


@pytest.mark.asyncio
@pytest.mark.parametrize(
    ("run", "expected"),
    [
        (service.AgentRun(stop_reason="timeout"), "timeout-no-output"),
        (service.AgentRun(stop_reason="agy-auth-required"), "agy-auth-required"),
        (service.AgentRun(status="ERROR", exit_code=1), "agy-error:ERROR"),
        (service.AgentRun(exit_code=3), "agy-exit:3"),
        (service.AgentRun(status="SUCCESS", exit_code=0), "no-output"),
    ],
)
async def test_missing_output_keeps_the_original_with_a_reason(tmp_path, run, expected):
    original = _deck()
    assert await service.judge_output(tmp_path, original, run, "soffice") == (
        original,
        "unchanged",
        expected,
    )


@pytest.mark.asyncio
async def test_whatever_opens_and_renders_is_delivered(tmp_path, monkeypatch):
    original = _deck()
    # No citation left at all: the agent is allowed to drop them.
    rewritten = _deck("Kısa ve atıfsız bir başlık")
    render = AsyncMock(return_value=True)
    monkeypatch.setattr(service, "validate_with_soffice", render)
    (tmp_path / service.OUTPUT_NAME).write_bytes(rewritten)

    finished = service.AgentRun(status="SUCCESS", exit_code=0)
    assert await service.judge_output(tmp_path, original, finished, "soffice") == (
        rewritten,
        "polished",
        "agy-SUCCESS",
    )
    assert render.await_args.kwargs["expected_pages"] == 1

    # Saved before the time ran out: still the agent's deck.
    stopped = service.AgentRun(stop_reason="timeout")
    assert await service.judge_output(tmp_path, original, stopped, "soffice") == (
        rewritten,
        "polished",
        "timeout-partial",
    )


@pytest.mark.asyncio
@pytest.mark.parametrize(
    ("output", "renders", "expected"),
    [
        ("same", True, "agent-made-no-changes"),
        (b"half-written", True, "invalid-pptx"),
        (None, False, "render-failed"),
    ],
)
async def test_output_that_is_not_a_usable_deck_is_rejected(tmp_path, monkeypatch, output, renders, expected):
    original = _deck()
    data = original if output == "same" else output or _deck("Çizilemeyen sunum")
    (tmp_path / service.OUTPUT_NAME).write_bytes(data)
    monkeypatch.setattr(service, "validate_with_soffice", AsyncMock(return_value=renders))
    monkeypatch.setattr(service, "MAX_REQUEST_BYTES", 10 * 1024 * 1024)
    result = await service.judge_output(
        tmp_path, original, service.AgentRun(status="SUCCESS", exit_code=0), "soffice"
    )
    assert result == (original, "unchanged", expected)


async def _never_disconnected() -> bool:
    return False


async def _run_fake_agent(tmp_path, code: str, **kwargs) -> service.AgentRun:
    options = {"timeout_s": 30.0, "is_disconnected": _never_disconnected, **kwargs}
    return await service.run_agent(
        [sys.executable, "-c", code],
        workdir=tmp_path,
        env=service._agent_environment(),
        **options,
    )


@pytest.mark.asyncio
async def test_agent_envelope_is_read(tmp_path, monkeypatch):
    monkeypatch.setattr(service, "_POLL_S", 0.05)
    run = await _run_fake_agent(
        tmp_path,
        "import json; print('log'); "
        "print(json.dumps({'status': 'SUCCESS', 'num_turns': 7, 'response': 'Başlıklar kısaldı.'}))",
    )
    assert (run.status, run.turns, run.exit_code, run.stop_reason) == ("SUCCESS", 7, 0, "")
    assert run.summary == "Başlıklar kısaldı."


@pytest.mark.asyncio
async def test_success_at_the_print_timeout_is_a_timeout(tmp_path, monkeypatch):
    """agy said SUCCESS after "print timeout after 10m0s with turn in progress"."""
    monkeypatch.setattr(service, "_POLL_S", 0.05)
    run = await _run_fake_agent(
        tmp_path,
        "import json, time; time.sleep(1.5); "
        "print(json.dumps({'status': 'SUCCESS', 'num_turns': 1, 'conversation_id': 'c-1'}))",
        timeout_s=2.0,
    )
    assert (run.status, run.stop_reason, run.conversation_id) == ("SUCCESS", "timeout", "c-1")


@pytest.mark.asyncio
async def test_login_prompt_stops_the_agent_at_once(tmp_path, monkeypatch):
    monkeypatch.setattr(service, "_POLL_S", 0.05)
    run = await _run_fake_agent(
        tmp_path,
        "import sys, time; "
        "print('Authentication required. Please visit the URL to log in:', file=sys.stderr, flush=True); "
        "time.sleep(60)",
    )
    assert run.stop_reason == "agy-auth-required"
    assert run.exit_code is not None
    assert run.duration_s < 20


@pytest.mark.asyncio
async def test_time_limit_and_hang_up_stop_the_agent(tmp_path, monkeypatch):
    monkeypatch.setattr(service, "_POLL_S", 0.05)
    monkeypatch.setattr(service, "_AGENT_KILL_GRACE_S", 0.0)
    sleeper = "import time; time.sleep(60)"

    timed_out = await _run_fake_agent(tmp_path, sleeper, timeout_s=0.3)
    assert timed_out.stop_reason == "timeout"
    assert timed_out.exit_code is not None
    assert timed_out.duration_s < 20

    async def gone() -> bool:
        return True

    hung_up = await _run_fake_agent(tmp_path, sleeper, is_disconnected=gone)
    assert hung_up.stop_reason == "client-disconnected"
    assert hung_up.exit_code is not None


@pytest.mark.asyncio
async def test_stop_process_terminates_and_reaps_child():
    proc = await asyncio.create_subprocess_exec(
        sys.executable,
        "-c",
        "import time; time.sleep(60)",
        stdout=asyncio.subprocess.DEVNULL,
        stderr=asyncio.subprocess.DEVNULL,
    )
    await service._stop_process(proc)
    assert proc.returncode is not None


def test_render_helper_reports_missing_inputs(tmp_path, monkeypatch):
    assert render_helper.main([str(tmp_path / "missing.pptx")]) == 2
    deck = tmp_path / "cilali.pptx"
    deck.write_bytes(_deck())
    monkeypatch.setattr(render_helper, "find_soffice", lambda: None)
    assert render_helper.main([str(deck), "--out", str(tmp_path / "render")]) == 1


def test_outline_lists_boxes_and_flags_collisions(tmp_path, capsys):
    """Slide 9 of the first polished deck kept an old text box under the new body text."""
    prs = pptx.Presentation()

    def box(slide, name, left, top, width, height, text):
        shape = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
        shape.name = name
        shape.text_frame.text = text

    crowded = prs.slides.add_slide(prs.slide_layouts[6])
    box(crowded, "sections[].consensus", 1, 1, 6, 3, "Ana metin: ğüşıöç")
    box(crowded, "sections[].implications_3", 1, 3, 6, 1, "Eski kutu")
    box(crowded, "rail_number", 0, 0, 1.2, 1.2, "3.1")  # template chrome, never flagged
    clean = prs.slides.add_slide(prs.slide_layouts[6])
    box(clean, "summary", 1, 1, 6, 3, "Tek kutu")
    wide = prs.slides.add_slide(prs.slide_layouts[6])
    box(wide, "wide", 8, 1, 3, 1, "Taşan kutu")
    deck = tmp_path / "cilali.pptx"
    prs.save(deck)

    result = outline_helper.outline(deck, slides=None, text_limit=5)
    first = result["slides"][0]
    assert [item["name"] for item in first["boxes"]] == [
        "sections[].consensus",
        "sections[].implications_3",
        "rail_number",
    ]
    assert first["boxes"][0]["text"] == "Ana m…"
    assert first["boxes"][0]["chars"] == len("Ana metin: ğüşıöç")
    assert first["issues"] == [
        "overlap: sections[].consensus <-> sections[].implications_3 (432 x 72 pt)"
    ]
    assert result["slides"][1]["issues"] == []
    assert result["slides"][2]["issues"] == ["off-slide: wide"]

    assert outline_helper.main([str(deck), "--check", "--slides", "1-3"]) == 0
    report = json.loads(capsys.readouterr().out)
    assert report["slides_with_problems"] == 2
    assert [problem["number"] for problem in report["problems"]] == [1, 3]
    assert outline_helper.main([str(tmp_path / "missing.pptx")]) == 2


# --------------------------------------------------------------------------- endpoint


@pytest.fixture
def endpoint(monkeypatch, tmp_path):
    from starlette.testclient import TestClient

    prompt_file = tmp_path / "istem.md"
    prompt_file.write_text("Sunumu toparla.", encoding="utf-8")
    monkeypatch.setenv("PRESENTATION_POLISHER_TOKEN", TOKEN)
    monkeypatch.setattr(service, "PROMPT_FILE", prompt_file)
    monkeypatch.setattr(service, "AGY_TIMEOUT_S", 600.0)
    monkeypatch.setattr(service, "REQUEST_BUDGET_S", 780.0)
    monkeypatch.setattr(service, "find_agy_binary", lambda: "agy")
    monkeypatch.setattr(service, "find_soffice_binary", lambda: "soffice")
    monkeypatch.setattr(service, "validate_with_soffice", AsyncMock(return_value=True))

    async def no_agent(*args, **kwargs):
        raise AssertionError("the agent must not be started")

    monkeypatch.setattr(service, "run_agent", no_agent)
    return TestClient(service.app)


def _post(client, content: bytes, token: str = TOKEN):
    return client.post(
        "/polish",
        content=content,
        headers={"Content-Type": PPTX_MEDIA_TYPE, "Authorization": f"Bearer {token}"},
        params={"language": "tr", "run_id": "run-1"},
    )


def test_endpoint_delivers_the_agents_deck(endpoint, monkeypatch):
    original, rewritten = _deck(), _deck("Ajanın sunumu")
    seen: dict = {}

    async def agent(command, *, workdir, env, timeout_s, is_disconnected):
        seen.update(
            command=command,
            files=sorted(path.name for path in workdir.iterdir()),
            input=(workdir / service.INPUT_NAME).read_bytes(),
            timeout=timeout_s,
        )
        (workdir / service.OUTPUT_NAME).write_bytes(rewritten)
        return service.AgentRun(status="SUCCESS", exit_code=0, duration_s=412.5, turns=31)

    monkeypatch.setattr(service, "run_agent", agent)
    response = _post(endpoint, original)

    assert response.status_code == 200
    assert response.content == rewritten
    assert response.headers["X-Presentation-Polisher-Status"] == "polished"
    assert response.headers["X-Presentation-Polisher-Reason"] == "agy-SUCCESS"
    assert response.headers["X-Presentation-Polisher-Agy-Status"] == "SUCCESS"
    assert response.headers["X-Presentation-Polisher-Duration-S"] == "412.5"
    assert response.headers["X-Presentation-Polisher-Turns"] == "31"
    assert seen["files"] == ["outline.py", "render", "render.py", "sunum.pptx"]
    assert seen["input"] == original
    assert seen["timeout"] == pytest.approx(600.0)
    assert seen["command"][-2] == "-p"
    assert seen["command"][-1].startswith("Sunumu toparla.")
    assert "cilali.pptx" in seen["command"][-1]


def test_endpoint_answers_with_the_original_when_the_agent_saves_nothing(endpoint, monkeypatch):
    original = _deck()

    async def agent(command, **kwargs):
        return service.AgentRun(stop_reason="timeout", duration_s=630.0)

    monkeypatch.setattr(service, "run_agent", agent)
    response = _post(endpoint, original)
    assert response.status_code == 200
    assert response.content == original
    assert response.headers["X-Presentation-Polisher-Status"] == "unchanged"
    assert response.headers["X-Presentation-Polisher-Reason"] == "timeout-no-output"


def test_a_long_queue_wait_skips_the_agent(endpoint, monkeypatch):
    monkeypatch.setattr(service, "REQUEST_BUDGET_S", 100.0)
    original = _deck()
    response = _post(endpoint, original)
    assert response.status_code == 200
    assert response.content == original
    assert response.headers["X-Presentation-Polisher-Reason"] == "busy"


def test_a_missing_prompt_file_skips_the_agent(endpoint, monkeypatch, tmp_path):
    monkeypatch.setattr(service, "PROMPT_FILE", tmp_path / "missing.md")
    response = _post(endpoint, _deck())
    assert response.status_code == 200
    assert response.headers["X-Presentation-Polisher-Status"] == "unchanged"
    assert response.headers["X-Presentation-Polisher-Reason"] == "prompt-missing"


def test_a_broken_word_list_skips_the_agent_and_shows_in_health(endpoint, monkeypatch, tmp_path):
    terms_file = tmp_path / "kelimeler.json"
    terms_file.write_text('{"tr": {"figür": ', encoding="utf-8")
    monkeypatch.setattr(service, "TERMS_FILE", terms_file)
    original = _deck()
    response = _post(endpoint, original)
    assert response.status_code == 200
    assert response.content == original
    assert response.headers["X-Presentation-Polisher-Status"] == "unchanged"
    assert response.headers["X-Presentation-Polisher-Reason"] == "terms-invalid"

    health = endpoint.get("/health")
    assert health.status_code == 503
    assert health.json()["terms"] is False


def test_endpoint_rejects_bad_requests(endpoint, monkeypatch):
    unauthorized = endpoint.post(
        "/polish", content=_deck(), headers={"Content-Type": PPTX_MEDIA_TYPE}
    )
    assert unauthorized.status_code == 401
    assert _post(endpoint, _deck(), token="wrong").status_code == 401
    assert _post(endpoint, b"not a presentation").status_code == 400

    monkeypatch.setattr(service, "MAX_REQUEST_BYTES", 4)
    assert _post(endpoint, b"12345").status_code == 413


def test_shared_service_token_does_not_open_the_service(endpoint, monkeypatch):
    monkeypatch.delenv("PRESENTATION_POLISHER_TOKEN")
    monkeypatch.setenv("SERVICE_TOKEN", TOKEN)
    assert _post(endpoint, _deck()).status_code == 503
    health = endpoint.get("/health")
    assert health.status_code == 503
    assert health.json()["authentication"] is False


def test_health_reports_readiness(endpoint, monkeypatch):
    ready = endpoint.get("/health")
    assert ready.status_code == 200
    body = ready.json()
    assert body["status"] == "ok"
    assert body["prompt"] is True and body["render_helper"] is True
    assert body["terms"] is True
    assert body["agent_timeout_s"] == 600.0
    assert "agy_path" not in body

    monkeypatch.setattr(service, "find_agy_binary", lambda: None)
    monkeypatch.setattr(service, "find_soffice_binary", lambda: None)
    degraded = endpoint.get("/health")
    assert degraded.status_code == 503
    assert degraded.json()["status"] == "degraded"
